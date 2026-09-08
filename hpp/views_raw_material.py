from django.core.paginator import Paginator
from django.db import transaction
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.db.models import Q, Sum, F, Count
from django.http import HttpResponse
from django.utils import timezone
from datetime import datetime, time
from decimal import Decimal
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from .models import UnitMaster, RawMaterial, RawMaterialUnitConversion, RawMaterialStockMutation

def _reconcile_material_stock(rm):
    """
    Rekonsiliasi kronologis saldo berjalan (balance_after) seluruh mutasi bahan baku
    dan update rm.current_stock agar 100% konsisten dengan riwayat kartu stok.
    """
    running_balance = Decimal(0)
    for mut in rm.mutations.all().order_by("created_at", "id"):
        if mut.mutation_type == "IN":
            running_balance += mut.stock_qty
        else:
            running_balance -= mut.stock_qty

        if mut.balance_after != running_balance:
            mut.balance_after = running_balance
            mut.save(update_fields=["balance_after"])

    if rm.current_stock != running_balance:
        rm.current_stock = running_balance
        rm.save(update_fields=["current_stock"])
    return running_balance


def raw_material_list(request):
    """
    Laporan Persediaan Bahan Baku (Valuasi Berdasarkan Harga Beli Terakhir)
    """
    query = request.GET.get("q", "").strip()
    category = request.GET.get("category", "").strip()
    low_stock = request.GET.get("low_stock", "").strip()

    materials_qs = RawMaterial.objects.all().prefetch_related("conversions").order_by("name")
    if query:
        materials_qs = materials_qs.filter(
            Q(name__icontains=query) |
            Q(code__icontains=query) |
            Q(category__icontains=query) |
            Q(notes__icontains=query)
        )
    if category:
        materials_qs = materials_qs.filter(category=category)
    if low_stock == "1":
        materials_qs = materials_qs.filter(current_stock__lte=F("minimum_stock"))

    # Hitung metrik KPI global
    base_qs = RawMaterial.objects.all()
    if category:
        base_qs = base_qs.filter(category=category)
    if query:
        base_qs = base_qs.filter(
            Q(name__icontains=query) |
            Q(code__icontains=query) |
            Q(category__icontains=query) |
            Q(notes__icontains=query)
        )
    total_inventory_value = sum(m.total_inventory_value for m in base_qs)
    total_sku_count = base_qs.count()
    low_stock_count = base_qs.filter(current_stock__lte=F("minimum_stock")).count()

    categories = RawMaterial.objects.values_list("category", flat=True).distinct()
    unit_materials = UnitMaster.objects.filter(category="raw_material", is_active=True).order_by("name")

    # Pagination 20 data per halaman
    paginator = Paginator(materials_qs, 20)
    page_number = request.GET.get("page", 1)
    materials = paginator.get_page(page_number)

    return render(request, "hpp/raw_material_list.html", {
        "materials": materials,
        "query": query,
        "selected_category": category,
        "low_stock": low_stock,
        "categories": categories,
        "total_inventory_value": total_inventory_value,
        "total_sku_count": total_sku_count,
        "low_stock_count": low_stock_count,
        "unit_materials": unit_materials,
    })


def raw_material_export_excel(request):
    """
    Export Katalog Persediaan Master Bahan Baku ke format Excel (.xlsx)
    """
    query = request.GET.get("q", "").strip()
    category = request.GET.get("category", "").strip()
    low_stock = request.GET.get("low_stock", "").strip()

    materials_qs = RawMaterial.objects.all().prefetch_related("conversions").order_by("name")
    if query:
        materials_qs = materials_qs.filter(
            Q(name__icontains=query) |
            Q(code__icontains=query) |
            Q(category__icontains=query) |
            Q(notes__icontains=query)
        )
    if category:
        materials_qs = materials_qs.filter(category=category)
    if low_stock == "1":
        materials_qs = materials_qs.filter(current_stock__lte=F("minimum_stock"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master Bahan Baku"

    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True)
    
    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    sec_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    warn_fill = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")
    
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    ws.merge_cells("A1:K1")
    ws["A1"] = "KATALOG & LAPORAN PERSEDIAAN BAHAN BAKU"
    ws["A1"].font = title_font
    ws.row_dimensions[1].height = 24

    ws["A2"] = f"Dicetak: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Kategori: {category or 'Semua'} | Filter Status: {'Hanya Kritis' if low_stock == '1' else 'Semua Stok'}"
    ws["A2"].font = subtitle_font
    ws.row_dimensions[2].height = 18

    ws.append([])

    headers = [
        "No", "Kode", "Nama Bahan Baku", "Kategori", "Satuan Dasar", 
        "Stok Tersedia", "Batas Minimum", "Status Stok", "Harga Beli Terakhir (Rp)", 
        "Total Nilai Persediaan (Rp)", "Catatan"
    ]
    ws.append(headers)
    h_idx = ws.max_row
    ws.row_dimensions[h_idx].height = 24

    for c_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=h_idx, column=c_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    total_value_all = Decimal(0)
    for idx, m in enumerate(materials_qs, 1):
        is_low = m.current_stock <= m.minimum_stock
        status_text = "KRITIS / RENDAH" if is_low else "AMAN"
        val = m.total_inventory_value
        total_value_all += val

        ws.append([
            idx,
            m.code,
            m.name,
            m.category or "Umum",
            m.stock_unit,
            float(m.current_stock),
            float(m.minimum_stock),
            status_text,
            float(m.last_purchase_price),
            float(val),
            m.notes or "-"
        ])
        r_idx = ws.max_row
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.border = thin_border
            if c_idx in [1, 2, 4, 5, 8]:
                cell.alignment = Alignment(horizontal="center")
            elif c_idx in [6, 7, 9, 10]:
                cell.alignment = Alignment(horizontal="right")
            if is_low and c_idx in [6, 8]:
                cell.fill = warn_fill
                cell.font = Font(name="Calibri", size=10, bold=True, color="9F1239")

    # Summary Row
    ws.append(["TOTAL NILAI INVENTORI", "", "", "", "", "", "", "", "", float(total_value_all), ""])
    sum_idx = ws.max_row
    for c_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=sum_idx, column=c_idx)
        cell.font = bold_font
        cell.fill = sec_fill
        cell.border = thin_border

    # Auto Column Widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len and '\n' not in val_str:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 40)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"Katalog_Bahan_Baku_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@transaction.atomic
def raw_material_create(request):
    """
    Tambah Master Bahan Baku Baru beserta Multi-Satuan (dengan validasi nama unik & transaksi atomic)
    """
    if request.method == "POST":
        code = request.POST.get("code", "").strip()
        name = request.POST.get("name", "").strip()
        category = request.POST.get("category", "Umum").strip()
        stock_unit = request.POST.get("stock_unit", "pcs").strip()
        initial_stock = Decimal(request.POST.get("initial_stock") or 0)
        last_purchase_price = Decimal(request.POST.get("last_purchase_price") or 0)
        minimum_stock = Decimal(request.POST.get("minimum_stock") or 0)
        notes = request.POST.get("notes", "").strip()

        if not name:
            messages.error(request, "Nama bahan baku wajib diisi.")
            return redirect("raw_material_list")

        if RawMaterial.objects.filter(name__iexact=name).exists():
            messages.error(request, f"Bahan baku dengan nama '{name}' sudah terdaftar dalam sistem.")
            return redirect("raw_material_list")

        if not code:
            last = RawMaterial.objects.order_by("-id").first()
            next_id = (last.id + 1) if last else 1
            code = f"MAT-{next_id:04d}"

        if RawMaterial.objects.filter(code__iexact=code).exists():
            messages.error(request, f"Kode '{code}' sudah digunakan oleh bahan baku lain.")
            return redirect("raw_material_list")

        rm = RawMaterial.objects.create(
            code=code,
            name=name,
            category=category,
            stock_unit=stock_unit,
            current_stock=initial_stock,
            last_purchase_price=last_purchase_price,
            minimum_stock=minimum_stock,
            notes=notes,
        )

        if initial_stock > 0:
            RawMaterialStockMutation.objects.create(
                raw_material=rm,
                mutation_type="IN",
                input_qty=initial_stock,
                input_unit=stock_unit,
                stock_qty=initial_stock,
                unit_price=last_purchase_price,
                total_price=initial_stock * last_purchase_price,
                balance_after=initial_stock,
                reference_no=f"INIT-{rm.code}",
                notes="Saldo stock awal sistem"
            )

        # Multi-Satuan
        conv_unit_names = request.POST.getlist("conv_unit_names[]")
        conv_factors = request.POST.getlist("conv_factors[]")
        conv_notes_list = request.POST.getlist("conv_notes[]")

        for u_name, factor_str, c_note in zip(conv_unit_names, conv_factors, conv_notes_list):
            u_name = u_name.strip()
            if u_name and u_name.lower() != stock_unit.lower():
                try:
                    c_factor = Decimal(factor_str or 1)
                    if c_factor > 0:
                        RawMaterialUnitConversion.objects.update_or_create(
                            raw_material=rm,
                            unit_name=u_name,
                            defaults={
                                "conversion_factor": c_factor,
                                "notes": c_note.strip()
                            }
                        )
                except Exception:
                    pass

        messages.success(request, f"Bahan baku '{name}' berhasil didaftarkan.")
        return redirect("raw_material_list")

    return render(request, "hpp/raw_material_form.html")


@transaction.atomic
def raw_material_update(request, uuid):
    rm = get_object_or_404(RawMaterial, uuid=uuid)
    if request.method == "POST":
        code = request.POST.get("code", rm.code).strip()
        name = request.POST.get("name", rm.name).strip()
        category = request.POST.get("category", rm.category).strip()
        stock_unit = request.POST.get("stock_unit", rm.stock_unit).strip()
        last_purchase_price = Decimal(request.POST.get("last_purchase_price") or 0)
        minimum_stock = Decimal(request.POST.get("minimum_stock") or 0)
        notes = request.POST.get("notes", "").strip()

        if not name:
            messages.error(request, "Nama bahan baku tidak boleh kosong.")
            return redirect("raw_material_detail", uuid=rm.uuid)

        if RawMaterial.objects.filter(name__iexact=name).exclude(uuid=rm.uuid).exists():
            messages.error(request, f"Nama bahan baku '{name}' sudah digunakan oleh item lain.")
            return redirect("raw_material_detail", uuid=rm.uuid)

        if code and RawMaterial.objects.filter(code__iexact=code).exclude(uuid=rm.uuid).exists():
            messages.error(request, f"Kode bahan baku '{code}' sudah digunakan oleh item lain.")
            return redirect("raw_material_detail", uuid=rm.uuid)

        # Proteksi Satuan Dasar
        if stock_unit.lower() != rm.stock_unit.lower():
            if rm.mutations.exists():
                messages.error(
                    request,
                    f"Satuan dasar persediaan '{rm.stock_unit}' tidak dapat diubah karena bahan baku ini sudah memiliki riwayat transaksi mutasi. "
                    "Untuk satuan lain, silakan daftarkan melalui fitur Multi-Satuan Konversi."
                )
                stock_unit = rm.stock_unit
            else:
                rm.stock_unit = stock_unit
        else:
            rm.stock_unit = stock_unit

        rm.code = code or rm.code
        rm.name = name
        rm.category = category
        rm.last_purchase_price = last_purchase_price
        rm.minimum_stock = minimum_stock
        rm.notes = notes
        rm.save()

        # Sinkronisasi Multi Satuan Konversi
        conv_unit_names = request.POST.getlist("conv_unit_names[]")
        conv_factors = request.POST.getlist("conv_factors[]")
        conv_notes_list = request.POST.getlist("conv_notes[]")

        submitted_units = set()
        for u_name, factor_str, c_note in zip(conv_unit_names, conv_factors, conv_notes_list):
            u_name = u_name.strip()
            if u_name and u_name.lower() != rm.stock_unit.lower():
                try:
                    c_factor = Decimal(factor_str or 1)
                    if c_factor > 0:
                        submitted_units.add(u_name.lower())
                        RawMaterialUnitConversion.objects.update_or_create(
                            raw_material=rm,
                            unit_name=u_name,
                            defaults={
                                "conversion_factor": c_factor,
                                "notes": c_note.strip()
                            }
                        )
                except Exception:
                    pass

        for old_conv in rm.conversions.all():
            if old_conv.unit_name.lower() not in submitted_units:
                is_used = RawMaterialStockMutation.objects.filter(raw_material=rm, input_unit__iexact=old_conv.unit_name).exists()
                if is_used:
                    messages.warning(request, f"Satuan '{old_conv.unit_name}' tetap dipertahankan karena sudah pernah digunakan dalam riwayat transaksi.")
                else:
                    old_conv.delete()

        messages.success(request, f"Data bahan baku '{rm.name}' berhasil diperbarui.")
        next_url = request.POST.get("next")
        if next_url:
            return redirect(next_url)
        return redirect("raw_material_detail", uuid=rm.uuid)
    return redirect("raw_material_detail", uuid=rm.uuid)

def raw_material_delete(request, uuid):
    rm = get_object_or_404(RawMaterial, uuid=uuid)
    if request.method == "POST":
        if rm.mutations.exists():
            messages.error(request, f"Bahan baku '{rm.name}' tidak dapat dihapus karena sudah memiliki riwayat mutasi.")
        else:
            name = rm.name
            rm.delete()
            messages.success(request, f"Bahan baku '{name}' berhasil dihapus.")
    else:
        messages.warning(request, "Penghapusan bahan baku harus dilakukan melalui tombol yang tersedia.")
    return redirect("raw_material_list")

def raw_material_conversion_add(request, material_uuid):
    rm = get_object_or_404(RawMaterial, uuid=material_uuid)
    if request.method == "POST":
        unit_name = request.POST.get("unit_name", "").strip()
        conversion_factor = Decimal(request.POST.get("conversion_factor") or 1)
        notes = request.POST.get("notes", "").strip()

        if conversion_factor <= 0:
            messages.error(request, "Faktor konversi harus lebih besar dari 0.")
            return redirect("raw_material_list")

        RawMaterialUnitConversion.objects.update_or_create(
            raw_material=rm,
            unit_name=unit_name,
            defaults={
                "conversion_factor": conversion_factor,
                "notes": notes
            }
        )
        messages.success(request, f"Konversi 1 {unit_name} = {conversion_factor} {rm.stock_unit} berhasil disimpan.")
    return redirect("raw_material_list")

def raw_material_conversion_delete(request, uuid):
    conv = get_object_or_404(RawMaterialUnitConversion, uuid=uuid)
    rm = conv.raw_material

    # Cek apakah satuan konversi ini pernah digunakan dalam transaksi mutasi
    if RawMaterialStockMutation.objects.filter(raw_material=rm, input_unit__iexact=conv.unit_name).exists():
        messages.error(request, f"Satuan konversi '{conv.unit_name}' tidak dapat dihapus karena sudah tercatat dalam riwayat transaksi mutasi stok!")
    else:
        unit_name = conv.unit_name
        conv.delete()
        messages.success(request, f"Konversi satuan '{unit_name}' berhasil dihapus.")

    next_url = request.GET.get("next")
    if next_url:
        return redirect(next_url)
    return redirect("raw_material_detail", uuid=rm.uuid)

def raw_material_stock_mutation_list(request):
    mat_uuid = request.GET.get("mat", "").strip()
    mutation_type = request.GET.get("type", "").strip()
    query = request.GET.get("q", "").strip()
    start_date_str = request.GET.get("start_date", "").strip()
    end_date_str = request.GET.get("end_date", "").strip()

    mutations_qs = RawMaterialStockMutation.objects.all().select_related("raw_material", "project")

    if mat_uuid:
        mutations_qs = mutations_qs.filter(raw_material__uuid=mat_uuid)
    if mutation_type:
        mutations_qs = mutations_qs.filter(mutation_type=mutation_type)
    if query:
        mutations_qs = mutations_qs.filter(
            Q(reference_no__icontains=query) |
            Q(notes__icontains=query) |
            Q(raw_material__name__icontains=query) |
            Q(raw_material__code__icontains=query)
        )

    # Filter rentang tanggal
    if start_date_str:
        try:
            start_date = datetime.strptime(start_date_str, "%Y-%m-%d").date()
            start_datetime = timezone.make_aware(datetime.combine(start_date, time.min))
            mutations_qs = mutations_qs.filter(created_at__gte=start_datetime)
        except ValueError:
            start_date_str = ""

    if end_date_str:
        try:
            end_date = datetime.strptime(end_date_str, "%Y-%m-%d").date()
            end_datetime = timezone.make_aware(datetime.combine(end_date, time.max))
            mutations_qs = mutations_qs.filter(created_at__lte=end_datetime)
        except ValueError:
            end_date_str = ""

    # Statistik ringkasan untuk hasil filter
    stats = mutations_qs.aggregate(
        total_in=Count("id", filter=Q(mutation_type="IN")),
        total_out=Count("id", filter=Q(mutation_type="OUT")),
        total_val=Sum("total_price")
    )
    total_in_count = stats["total_in"] or 0
    total_out_count = stats["total_out"] or 0
    total_val = stats["total_val"] or Decimal(0)
    total_count = total_in_count + total_out_count

    materials = RawMaterial.objects.all().prefetch_related("conversions").order_by("name")

    selected_mat_obj = None
    if mat_uuid:
        selected_mat_obj = RawMaterial.objects.filter(uuid=mat_uuid).first()

    materials_json = [
        {
            "uuid": str(m.uuid),
            "code": m.code,
            "name": m.name,
            "stock_unit": m.stock_unit,
            "current_stock": float(m.current_stock),
            "last_price": float(m.last_purchase_price),
            "conversions": [
                {
                    "unit_name": conv.unit_name,
                    "factor": float(conv.conversion_factor)
                }
                for conv in m.conversions.all()
            ]
        }
        for m in materials
    ]

    # Pagination 20 data per halaman
    paginator = Paginator(mutations_qs, 20)
    page_number = request.GET.get("page", 1)
    mutations = paginator.get_page(page_number)

    return render(request, "hpp/raw_material_mutation_list.html", {
        "mutations": mutations,
        "materials": materials,
        "materials_json": json.dumps(materials_json),
        "selected_mat": mat_uuid,
        "selected_mat_obj": selected_mat_obj,
        "selected_type": mutation_type,
        "query": query,
        "start_date_str": start_date_str,
        "end_date_str": end_date_str,
        "total_count": total_count,
        "total_in_count": total_in_count,
        "total_out_count": total_out_count,
        "total_val": total_val,
    })


def raw_material_mutation_export_excel(request):
    """
    Export Rekapitulasi Log Mutasi Stok Bahan Baku ke File Excel (.xlsx)
    Sesuai filter pencarian, bahan baku, tipe mutasi, dan periode tanggal.
    """
    mat_uuid = request.GET.get("mat", "").strip()
    mutation_type = request.GET.get("type", "").strip()
    query = request.GET.get("q", "").strip()
    start_date_str = request.GET.get("start_date", "").strip()
    end_date_str = request.GET.get("end_date", "").strip()

    mutations_qs = RawMaterialStockMutation.objects.all().select_related("raw_material", "project")

    selected_mat = None
    if mat_uuid:
        selected_mat = RawMaterial.objects.filter(uuid=mat_uuid).first()
        if selected_mat:
            mutations_qs = mutations_qs.filter(raw_material=selected_mat)
    if mutation_type:
        mutations_qs = mutations_qs.filter(mutation_type=mutation_type)
    if query:
        mutations_qs = mutations_qs.filter(
            Q(reference_no__icontains=query) |
            Q(notes__icontains=query) |
            Q(raw_material__name__icontains=query) |
            Q(raw_material__code__icontains=query)
        )

    if start_date_str:
        try:
            start_date = datetime.strptime(start_date_str, "%Y-%m-%d").date()
            start_datetime = timezone.make_aware(datetime.combine(start_date, time.min))
            mutations_qs = mutations_qs.filter(created_at__gte=start_datetime)
        except ValueError:
            start_date_str = ""

    if end_date_str:
        try:
            end_date = datetime.strptime(end_date_str, "%Y-%m-%d").date()
            end_datetime = timezone.make_aware(datetime.combine(end_date, time.max))
            mutations_qs = mutations_qs.filter(created_at__lte=end_datetime)
        except ValueError:
            end_date_str = ""

    mutations_qs = mutations_qs.order_by("-created_at", "-id")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rekap Mutasi"

    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")  # Slate 800
    thin_border_side = Side(style="thin", color="CBD5E1")
    border_all = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    ws.merge_cells("A1:N1")
    ws["A1"] = "REKAPITULASI LOG MUTASI STOK BAHAN BAKU"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 25

    desc_parts = []
    if selected_mat:
        desc_parts.append(f"Material: {selected_mat.name} ({selected_mat.code})")
    if mutation_type:
        desc_parts.append(f"Tipe: {mutation_type}")
    if start_date_str and end_date_str:
        desc_parts.append(f"Periode: {start_date_str} s/d {end_date_str}")
    elif start_date_str:
        desc_parts.append(f"Sejak: {start_date_str}")
    elif end_date_str:
        desc_parts.append(f"Hingga: {end_date_str}")
    if query:
        desc_parts.append(f"Pencarian: '{query}'")
    if not desc_parts:
        desc_parts.append("Semua Riwayat Transaksi")

    ws["A2"] = " | ".join(desc_parts)
    ws["A2"].font = subtitle_font
    ws.row_dimensions[2].height = 18

    headers = [
        "No", "Tanggal & Waktu", "Kode Material", "Nama Bahan Baku", 
        "Arah", "Jumlah Input", "Satuan Input", "Jumlah Dasar", 
        "Satuan Dasar", "Saldo Akhir", "Harga Satuan (Rp)", "Total Nilai (Rp)", "No. Ref / Project", "Keterangan"
    ]
    ws.append([])
    ws.append(headers)
    header_row_idx = ws.max_row
    ws.row_dimensions[header_row_idx].height = 22

    for col_num in range(1, len(headers) + 1):
        c = ws.cell(row=header_row_idx, column=col_num)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_all

    for idx, mut in enumerate(mutations_qs, start=1):
        dt_str = timezone.localtime(mut.created_at).strftime("%Y-%m-%d %H:%M") if mut.created_at else "-"
        proj_ref = f"[{mut.project.code}] {mut.project.name}" if mut.project else (mut.reference_no or "-")
        row = [
            idx,
            dt_str,
            mut.raw_material.code,
            mut.raw_material.name,
            mut.mutation_type,
            float(mut.input_qty),
            mut.input_unit,
            float(mut.stock_qty),
            mut.raw_material.stock_unit,
            float(mut.balance_after),
            float(mut.unit_price),
            float(mut.total_price),
            proj_ref,
            mut.notes or "-"
        ]
        ws.append(row)
        curr_row = ws.max_row
        for c_idx in range(1, len(row) + 1):
            cell = ws.cell(row=curr_row, column=c_idx)
            cell.border = border_all
            if c_idx in [1, 2, 5, 7, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in [6, 8, 10, 11, 12]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if c_idx in [11, 12]:
                    cell.number_format = "#,##0.00"
                else:
                    cell.number_format = "#,##0.0000"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    col_widths = {
        "A": 6, "B": 18, "C": 14, "D": 28, "E": 8, "F": 14, "G": 14,
        "H": 14, "I": 14, "J": 14, "K": 18, "L": 20, "M": 24, "N": 30
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"Rekap_Mutasi_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

@transaction.atomic
def raw_material_mutation_create(request):
    """
    Input Mutasi Manual Masuk (Pembelian Supplier) atau Keluar (Rusak / Pemakaian Non-Project)
    """
    if request.method == "POST":
        mat_uuid = request.POST.get("material_uuid")
        mutation_type = request.POST.get("mutation_type", "IN")
        input_qty = Decimal(request.POST.get("input_qty") or 0)
        input_unit = request.POST.get("input_unit", "").strip()
        unit_price = Decimal(request.POST.get("unit_price") or 0)
        reference_no = request.POST.get("reference_no", "").strip()
        notes = request.POST.get("notes", "").strip()

        rm = get_object_or_404(RawMaterial, uuid=mat_uuid)

        if input_qty <= 0:
            messages.error(request, "Jumlah mutasi harus lebih besar dari 0.")
            return redirect("raw_material_stock_mutation_list")

        # Cek faktor konversi jika input_unit berbeda dari stock_unit
        factor = Decimal(1)
        if input_unit and input_unit.lower() != rm.stock_unit.lower():
            conv = rm.conversions.filter(unit_name__iexact=input_unit).first()
            if conv:
                factor = conv.conversion_factor
            else:
                input_unit = rm.stock_unit
        else:
            input_unit = rm.stock_unit

        base_qty = input_qty * factor
        total_price = input_qty * unit_price

        if mutation_type == "IN":
            rm.current_stock += base_qty
            if unit_price > 0:
                # Update harga beli terakhir per stock_unit
                rm.last_purchase_price = unit_price / factor
            prefix = "[MASUK/BELI]"
        else:
            if rm.current_stock < base_qty:
                messages.error(request, f"Stok {rm.name} tidak mencukupi (Tersedia: {rm.current_stock} {rm.stock_unit}).")
                return redirect("raw_material_stock_mutation_list")
            rm.current_stock -= base_qty
            prefix = "[RUSAK/OUT]"

        rm.save()

        full_notes = f"{prefix} {notes}".strip()
        RawMaterialStockMutation.objects.create(
            raw_material=rm,
            mutation_type=mutation_type,
            input_qty=input_qty,
            input_unit=input_unit,
            stock_qty=base_qty,
            unit_price=unit_price,
            total_price=total_price,
            balance_after=rm.current_stock,
            reference_no=reference_no or f"MUT-{rm.code}",
            notes=full_notes
        )
        _reconcile_material_stock(rm)
        messages.success(request, f"Mutasi bahan baku {rm.name} ({mutation_type} {input_qty} {input_unit}) berhasil dicatat.")
    return redirect("raw_material_stock_mutation_list")

def _get_stock_card_data(mat_uuid, start_date_str, end_date_str):
    """
    Helper untuk menghitung data kartu stok bahan baku:
    - selected_mat
    - beginning_balance
    - total_in (periode)
    - total_out (periode)
    - ending_balance (periode)
    - mutations_qs
    - start_date, end_date
    """
    materials = RawMaterial.objects.all().order_by("name").prefetch_related("conversions")
    if mat_uuid:
        selected_mat = RawMaterial.objects.filter(uuid=mat_uuid).prefetch_related("conversions").first()
    else:
        selected_mat = materials.first()

    start_date = None
    end_date = None
    start_datetime = None
    end_datetime = None

    if start_date_str:
        try:
            start_date = datetime.strptime(start_date_str, "%Y-%m-%d").date()
            start_datetime = timezone.make_aware(datetime.combine(start_date, time.min))
        except ValueError:
            start_date_str = ""

    if end_date_str:
        try:
            end_date = datetime.strptime(end_date_str, "%Y-%m-%d").date()
            end_datetime = timezone.make_aware(datetime.combine(end_date, time.max))
        except ValueError:
            end_date_str = ""

    beginning_balance = Decimal(0)
    total_in = Decimal(0)
    total_out = Decimal(0)
    ending_balance = Decimal(0)
    mutations_qs = RawMaterialStockMutation.objects.none()

    if selected_mat:
        # Hitung Saldo Awal sebelum rentang tanggal yang dipilih
        if start_datetime:
            prior_mutations = selected_mat.mutations.filter(created_at__lt=start_datetime)
            last_prior = prior_mutations.order_by("-created_at", "-id").first()
            if last_prior:
                beginning_balance = last_prior.balance_after
            else:
                beginning_balance = Decimal(0)
        else:
            beginning_balance = Decimal(0)

        # Mutasi dalam periode
        mutations_qs = selected_mat.mutations.all().select_related("project")
        if start_datetime:
            mutations_qs = mutations_qs.filter(created_at__gte=start_datetime)
        if end_datetime:
            mutations_qs = mutations_qs.filter(created_at__lte=end_datetime)

        totals = mutations_qs.aggregate(
            total_in=Sum("stock_qty", filter=Q(mutation_type="IN")),
            total_out=Sum("stock_qty", filter=Q(mutation_type="OUT"))
        )
        total_in = totals["total_in"] or Decimal(0)
        total_out = totals["total_out"] or Decimal(0)

        ending_balance = beginning_balance + total_in - total_out

    return {
        "materials": materials,
        "selected_mat": selected_mat,
        "start_date_str": start_date_str,
        "end_date_str": end_date_str,
        "start_date": start_date,
        "end_date": end_date,
        "beginning_balance": beginning_balance,
        "total_in": total_in,
        "total_out": total_out,
        "ending_balance": ending_balance,
        "mutations_qs": mutations_qs,
    }


def raw_material_stock_card(request, uuid):
    """
    Kartu Stok Bahan Baku Tunggal -> Redirect ke index dengan parameter ?mat=uuid
    agar selalu menikmati fitur lengkap (filter tanggal, export excel, print view, combobox)
    """
    return redirect(f"/raw-materials/cards/?mat={uuid}")


def raw_material_stock_card_index(request):
    """
    Kartu Stok Bahan Baku Standalone (Filter Pilih Item, Rentang Tanggal, Saldo Awal, Pagination)
    """
    mat_uuid = request.GET.get("mat", "").strip()
    start_date_str = request.GET.get("start_date", "").strip()
    end_date_str = request.GET.get("end_date", "").strip()
    page_number = request.GET.get("page", 1)

    data = _get_stock_card_data(mat_uuid, start_date_str, end_date_str)
    selected_mat = data["selected_mat"]
    mutations_qs = data["mutations_qs"]

    # Pagination 25 data per halaman
    paginator = Paginator(mutations_qs, 25)
    mutations = paginator.get_page(page_number)

    # Serialisasi data bahan baku untuk combobox Alpine.js
    materials_json = [
        {
            "uuid": str(m.uuid),
            "code": m.code,
            "name": m.name,
            "stock_unit": m.stock_unit,
            "current_stock": float(m.current_stock),
            "last_price": float(m.last_purchase_price)
        }
        for m in data["materials"]
    ]

    total_inventory_value = Decimal(0)
    if selected_mat:
        total_inventory_value = data["ending_balance"] * selected_mat.last_purchase_price

    return render(request, "hpp/raw_material_stock_card_index.html", {
        "materials": data["materials"],
        "materials_json": json.dumps(materials_json),
        "selected_mat": selected_mat,
        "mutations": mutations,
        "start_date": data["start_date_str"],
        "end_date": data["end_date_str"],
        "beginning_balance": data["beginning_balance"],
        "total_in": data["total_in"],
        "total_out": data["total_out"],
        "ending_balance": data["ending_balance"],
        "total_inventory_value": total_inventory_value,
    })


def raw_material_stock_card_export_excel(request):
    """
    Export Kartu Stok Bahan Baku ke format Excel (.xlsx) menggunakan openpyxl
    """
    mat_uuid = request.GET.get("mat", "").strip()
    start_date_str = request.GET.get("start_date", "").strip()
    end_date_str = request.GET.get("end_date", "").strip()

    data = _get_stock_card_data(mat_uuid, start_date_str, end_date_str)
    selected_mat = data["selected_mat"]
    if not selected_mat:
        messages.error(request, "Bahan baku tidak ditemukan.")
        return redirect("raw_material_stock_card_index")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kartu Stok"

    # Fonts & Styles
    title_font = Font(name="Calibri", size=15, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    bold_font = Font(name="Calibri", size=10, bold=True)
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    
    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")  # Teal 700
    section_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    beg_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")  # Amber 100
    
    thin_border_side = Side(style="thin", color="CBD5E1")
    border_all = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # Title
    ws.merge_cells("A1:M1")
    ws["A1"] = f"KARTU STOK BAHAN BAKU: {selected_mat.name.upper()}"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 25

    period_desc = "Semua Riwayat Transaksi"
    if data["start_date_str"] and data["end_date_str"]:
        period_desc = f"Periode: {data['start_date_str']} s/d {data['end_date_str']}"
    elif data["start_date_str"]:
        period_desc = f"Periode Mulai: {data['start_date_str']}"
    elif data["end_date_str"]:
        period_desc = f"Periode Sampai: {data['end_date_str']}"

    ws["A2"] = f"Kode: {selected_mat.code} | Satuan Dasar: {selected_mat.stock_unit} | {period_desc}"
    ws["A2"].font = subtitle_font
    ws.row_dimensions[2].height = 18

    # Ringkasan KPI
    ws.append([])
    ws.append(["RINGKASAN PERSEDIAAN PERIODE"])
    ws.cell(row=ws.max_row, column=1).font = bold_font
    ws.cell(row=ws.max_row, column=1).fill = section_fill

    ws.append(["Saldo Awal", float(data["beginning_balance"]), selected_mat.stock_unit, "", "Harga Beli Terakhir", float(selected_mat.last_purchase_price)])
    ws.append(["Total Masuk (IN)", float(data["total_in"]), selected_mat.stock_unit, "", "Total Valuasi Saldo Akhir", float(data["ending_balance"] * selected_mat.last_purchase_price)])
    ws.append(["Total Keluar (OUT)", float(data["total_out"]), selected_mat.stock_unit, "", "", ""])
    ws.append(["Saldo Akhir Periode", float(data["ending_balance"]), selected_mat.stock_unit, "", "", ""])

    for r in range(4, 9):
        ws.cell(row=r, column=1).font = bold_font
        ws.cell(row=r, column=5).font = bold_font

    ws.append([])

    # Table Header
    headers = [
        "No", "Waktu / Tanggal", "No. Referensi", "Project Terkait", "Arah Mutasi", 
        "Satuan Input", "Qty Input", "Masuk (IN)", "Keluar (OUT)", "Saldo Berjalan", 
        "Harga Satuan (Rp)", "Total Nilai (Rp)", "Keterangan"
    ]
    ws.append(headers)
    header_row_idx = ws.max_row
    ws.row_dimensions[header_row_idx].height = 24

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Baris Saldo Awal jika rentang tanggal difilter
    if data["start_date_str"]:
        ws.append([
            "", f"{data['start_date_str']} 00:00", "-", "-", "SALDO AWAL", 
            selected_mat.stock_unit, float(data["beginning_balance"]), "-", "-", 
            float(data["beginning_balance"]), float(selected_mat.last_purchase_price), 
            float(data["beginning_balance"] * selected_mat.last_purchase_price), "Saldo Awal Sebelum Periode"
        ])
        beg_row_idx = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=beg_row_idx, column=col_idx)
            cell.fill = beg_fill
            cell.font = bold_font
            cell.border = border_all

    # Data Rows
    row_num = 1
    for m in data["mutations_qs"]:
        proj_str = f"#{m.project.code} - {m.project.name}" if m.project else "-"
        in_qty = float(m.stock_qty) if m.mutation_type == "IN" else ""
        out_qty = float(m.stock_qty) if m.mutation_type == "OUT" else ""
        
        ws.append([
            row_num,
            m.created_at.strftime("%Y-%m-%d %H:%M"),
            m.reference_no or "-",
            proj_str,
            m.get_mutation_type_display(),
            m.input_unit,
            float(m.input_qty),
            in_qty,
            out_qty,
            float(m.balance_after),
            float(m.unit_price),
            float(m.total_price),
            m.notes or "-"
        ])
        curr_row = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.border = border_all
            if col_idx in [1, 5, 6]:
                cell.alignment = Alignment(horizontal="center")
            elif col_idx in [7, 8, 9, 10, 11, 12]:
                cell.alignment = Alignment(horizontal="right")
        row_num += 1

    # Auto Column Widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len and '\n' not in val_str:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 40)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    safe_code = "".join(c for c in selected_mat.code if c.isalnum() or c in ("-", "_"))
    filename = f"Kartu_Stok_{safe_code}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def raw_material_detail(request, uuid):
    """
    Halaman Detail Profil Bahan Baku & Konfigurasi Multi-Satuan
    """
    rm = get_object_or_404(RawMaterial.objects.prefetch_related("conversions", "mutations"), uuid=uuid)
    unit_materials = UnitMaster.objects.filter(category="raw_material", is_active=True).order_by("name")
    recent_mutations = rm.mutations.select_related("project")[:10]

    return render(request, "hpp/raw_material_detail.html", {
        "material": rm,
        "unit_materials": unit_materials,
        "recent_mutations": recent_mutations,
    })


def raw_material_mutation_detail(request, uuid):
    """
    Halaman Detail Mutasi Bahan Baku
    """
    mutation = get_object_or_404(RawMaterialStockMutation.objects.select_related("raw_material", "project"), uuid=uuid)
    unit_materials = UnitMaster.objects.filter(category="raw_material", is_active=True).order_by("name")

    return render(request, "hpp/raw_material_mutation_detail.html", {
        "mutation": mutation,
        "unit_materials": unit_materials,
    })


@transaction.atomic
def raw_material_mutation_update(request, uuid):
    """
    Update Mutasi Manual Bahan Baku (Koreksi Jumlah / Harga / Catatan dengan Rekonsiliasi Saldo Berjalan)
    """
    mutation = get_object_or_404(RawMaterialStockMutation.objects.select_related("raw_material", "project"), uuid=uuid)
    
    if mutation.project:
        messages.error(request, "Mutasi yang berasal dari realisasi project tidak dapat diedit langsung dari sini. Silakan edit melalui menu Realisasi Project.")
        return redirect("raw_material_mutation_detail", uuid=mutation.uuid)

    if request.method == "POST":
        rm = mutation.raw_material
        mutation_type = request.POST.get("mutation_type", mutation.mutation_type)
        input_qty = Decimal(request.POST.get("input_qty") or 0)
        input_unit = request.POST.get("input_unit", "").strip() or rm.stock_unit
        unit_price = Decimal(request.POST.get("unit_price") or 0)
        reference_no = request.POST.get("reference_no", "").strip()
        notes = request.POST.get("notes", "").strip()

        if input_qty <= 0:
            messages.error(request, "Jumlah mutasi harus lebih besar dari 0.")
            return redirect("raw_material_mutation_detail", uuid=mutation.uuid)

        # Hitung faktor konversi baru
        factor = Decimal(1)
        if input_unit and input_unit.lower() != rm.stock_unit.lower():
            conv = rm.conversions.filter(unit_name__iexact=input_unit).first()
            if conv:
                factor = conv.conversion_factor
            else:
                input_unit = rm.stock_unit
        else:
            input_unit = rm.stock_unit

        base_qty = input_qty * factor
        total_price = input_qty * unit_price

        mutation.mutation_type = mutation_type
        mutation.input_qty = input_qty
        mutation.input_unit = input_unit
        mutation.stock_qty = base_qty
        mutation.unit_price = unit_price
        mutation.total_price = total_price
        mutation.reference_no = reference_no
        mutation.notes = notes
        mutation.save()

        # Update harga beli terakhir jika mutasi IN
        if mutation_type == "IN" and unit_price > 0:
            rm.last_purchase_price = unit_price / factor
            rm.save(update_fields=["last_purchase_price"])

        # Jalankan rekonsiliasi kronologis saldo berjalan
        new_bal = _reconcile_material_stock(rm)
        if new_bal < 0:
            transaction.set_rollback(True)
            messages.error(request, f"Penyesuaian tidak dapat disimpan karena akan menyebabkan saldo stok menjadi minus ({new_bal} {rm.stock_unit}).")
            return redirect("raw_material_mutation_detail", uuid=mutation.uuid)

        messages.success(request, f"Transaksi mutasi {reference_no or mutation.id} berhasil diperbarui dan saldo kartu stok telah diselaraskan.")
        return redirect("raw_material_mutation_detail", uuid=mutation.uuid)

    return redirect("raw_material_mutation_detail", uuid=mutation.uuid)


@transaction.atomic
def raw_material_mutation_delete(request, uuid):
    """
    Hapus Mutasi Manual & Rekonsiliasi Saldo Stok
    """
    mutation = get_object_or_404(RawMaterialStockMutation.objects.select_related("raw_material", "project"), uuid=uuid)

    if mutation.project:
        messages.error(request, "Mutasi ini berasal dari Realisasi Project. Untuk membatalkan, hapus transaksi melalui menu Realisasi Project.")
        return redirect("raw_material_stock_mutation_list")

    rm = mutation.raw_material
    ref = mutation.reference_no or str(mutation.id)
    mutation.delete()

    new_bal = _reconcile_material_stock(rm)
    if new_bal < 0:
        transaction.set_rollback(True)
        messages.error(request, f"Tidak dapat menghapus mutasi masuk ini karena mutasi keluar yang sudah ada melebihi stok yang tersedia (menghasilkan saldo minus {new_bal} {rm.stock_unit}).")
        return redirect("raw_material_stock_mutation_list")

    messages.success(request, f"Transaksi mutasi '{ref}' berhasil dihapus dan seluruh saldo kartu stok '{rm.name}' telah diselaraskan.")
    return redirect("raw_material_stock_mutation_list")
