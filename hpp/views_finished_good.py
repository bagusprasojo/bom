import json
from datetime import datetime, time
from decimal import Decimal
from django.core.paginator import Paginator
from django.db import transaction
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.db.models import Q, Sum, F, Count, DecimalField, ExpressionWrapper
from django.http import HttpResponse
from django.utils import timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import UnitMaster, FinishedGood, FinishedGoodStockMutation


def finished_good_list(request):
    """
    Daftar Master & Laporan Persediaan Barang Jadi (Finished Goods)
    Lengkap dengan kartu KPI, filter kategori, filter stok kritis, pencarian, dan pagination.
    """
    query = request.GET.get("q", "").strip()
    category = request.GET.get("category", "").strip()
    low_stock = request.GET.get("low_stock", "").strip()

    base_qs = FinishedGood.objects.all()

    # Hitung metrik KPI secara keseluruhan
    total_sku_count = base_qs.count()
    total_stock_units = sum(fg.current_stock for fg in base_qs)
    total_inventory_value = sum(fg.total_inventory_value for fg in base_qs)
    low_stock_count = base_qs.filter(current_stock__lte=F("minimum_stock")).count()

    fg_qs = base_qs.order_by("name")
    if query:
        fg_qs = fg_qs.filter(
            Q(name__icontains=query) |
            Q(sku__icontains=query) |
            Q(category__icontains=query) |
            Q(notes__icontains=query)
        )
    if category:
        fg_qs = fg_qs.filter(category=category)
    if low_stock == "1":
        fg_qs = fg_qs.filter(current_stock__lte=F("minimum_stock"))

    # Rekomendasi auto-generate SKU berikutnya
    last_fg = FinishedGood.objects.order_by("-id").first()
    next_id = (last_fg.id + 1) if last_fg else 1
    suggested_sku = f"FG-{next_id:04d}"
    while FinishedGood.objects.filter(sku__iexact=suggested_sku).exists():
        next_id += 1
        suggested_sku = f"FG-{next_id:04d}"

    # Daftar kategori & satuan aktif
    categories = FinishedGood.objects.values_list("category", flat=True).distinct()
    unit_fgs = UnitMaster.objects.filter(category="finished_good", is_active=True).order_by("name")

    # Pagination 20 data per halaman
    paginator = Paginator(fg_qs, 20)
    page_number = request.GET.get("page", 1)
    finished_goods = paginator.get_page(page_number)

    # Flag has_transactions untuk proteksi delete & edit satuan
    for fg in finished_goods:
        fg.has_transactions = fg.mutations.exists() or fg.project_usages.exists()

    return render(request, "hpp/finished_good_list.html", {
        "finished_goods": finished_goods,
        "query": query,
        "selected_category": category,
        "low_stock": low_stock,
        "categories": categories,
        "total_sku_count": total_sku_count,
        "total_stock_units": total_stock_units,
        "total_inventory_value": total_inventory_value,
        "low_stock_count": low_stock_count,
        "unit_fgs": unit_fgs,
        "suggested_sku": suggested_sku,
    })


@transaction.atomic
def finished_good_create(request):
    """
    Pendaftaran Barang Jadi Baru (Modal Cepat atau Form)
    Dilengkapi auto-generate SKU, validasi atomik, dan inisialisasi mutasi saldo awal.
    """
    if request.method == "POST":
        sku = request.POST.get("sku", "").strip()
        name = request.POST.get("name", "").strip()
        category = request.POST.get("category", "Umum").strip() or "Umum"
        unit = request.POST.get("unit", "unit").strip() or "unit"
        standard_cost = Decimal(request.POST.get("standard_cost") or 0)
        initial_stock = Decimal(request.POST.get("initial_stock") or 0)
        minimum_stock = Decimal(request.POST.get("minimum_stock") or 0)
        notes = request.POST.get("notes", "").strip()

        if not name:
            messages.error(request, "Nama barang jadi wajib diisi.")
            return redirect("finished_good_list")

        if FinishedGood.objects.filter(name__iexact=name).exists():
            messages.error(request, f"Barang jadi dengan nama '{name}' sudah terdaftar dalam sistem.")
            return redirect("finished_good_list")

        # Auto-generate SKU jika dikosongkan pengguna
        if not sku:
            last_fg = FinishedGood.objects.order_by("-id").first()
            next_id = (last_fg.id + 1) if last_fg else 1
            sku = f"FG-{next_id:04d}"
            while FinishedGood.objects.filter(sku__iexact=sku).exists():
                next_id += 1
                sku = f"FG-{next_id:04d}"

        if FinishedGood.objects.filter(sku__iexact=sku).exists():
            messages.error(request, f"Kode SKU '{sku}' sudah digunakan oleh barang jadi lain.")
            return redirect("finished_good_list")

        fg = FinishedGood.objects.create(
            sku=sku,
            name=name,
            category=category,
            unit=unit,
            standard_cost=standard_cost,
            current_stock=initial_stock,
            minimum_stock=minimum_stock,
            notes=notes,
        )

        # Otomatis catat transaksi mutasi masuk (IN) Saldo Awal jika ada stok awal
        if initial_stock > 0:
            FinishedGoodStockMutation.objects.create(
                finished_good=fg,
                mutation_type="IN",
                quantity=initial_stock,
                balance_after=initial_stock,
                reference_no=f"INIT-{fg.sku}",
                notes="Saldo stock awal barang jadi"
            )

        messages.success(request, f"Barang Jadi '{name}' ({sku}) berhasil didaftarkan ke master.")
        return redirect("finished_good_list")

    return redirect("finished_good_list")


@transaction.atomic
def finished_good_update(request, uuid):
    """
    Update Master Data Barang Jadi
    Dengan proteksi keunikan nama/SKU dan proteksi satuan jika sudah memiliki riwayat mutasi.
    """
    fg = get_object_or_404(FinishedGood, uuid=uuid)

    if request.method == "POST":
        sku = request.POST.get("sku", "").strip() or fg.sku
        name = request.POST.get("name", "").strip()
        category = request.POST.get("category", "Umum").strip() or "Umum"
        unit = request.POST.get("unit", "").strip() or fg.unit
        standard_cost = Decimal(request.POST.get("standard_cost") or 0)
        minimum_stock = Decimal(request.POST.get("minimum_stock") or 0)
        notes = request.POST.get("notes", "").strip()

        if not name:
            messages.error(request, "Nama barang jadi tidak boleh kosong.")
            return redirect("finished_good_list")

        if FinishedGood.objects.filter(name__iexact=name).exclude(uuid=fg.uuid).exists():
            messages.error(request, f"Barang jadi dengan nama '{name}' sudah digunakan oleh item lain.")
            return redirect("finished_good_list")

        if FinishedGood.objects.filter(sku__iexact=sku).exclude(uuid=fg.uuid).exists():
            messages.error(request, f"Kode SKU '{sku}' sudah digunakan oleh barang jadi lain.")
            return redirect("finished_good_list")

        # Proteksi satuan jika sudah ada transaksi mutasi atau pemakaian project
        has_mutations = fg.mutations.exists() or fg.project_usages.exists()
        if has_mutations and unit.lower() != fg.unit.lower():
            messages.warning(
                request,
                f"Satuan dasar '{fg.unit}' tidak dapat diubah karena barang jadi ini sudah memiliki riwayat transaksi mutasi stok. Perubahan lainnya berhasil disimpan."
            )
            unit = fg.unit
        else:
            fg.unit = unit

        fg.sku = sku
        fg.name = name
        fg.category = category
        fg.standard_cost = standard_cost
        fg.minimum_stock = minimum_stock
        fg.notes = notes
        fg.save()

        messages.success(request, f"Master barang jadi '{fg.name}' ({fg.sku}) berhasil diperbarui.")
        return redirect("finished_good_list")

    return redirect("finished_good_list")


@transaction.atomic
def finished_good_delete(request, uuid):
    """
    Hapus Master Barang Jadi (Hanya jika belum memiliki transaksi mutasi atau pemakaian project)
    """
    fg = get_object_or_404(FinishedGood, uuid=uuid)

    if fg.mutations.exists() or fg.project_usages.exists():
        messages.error(
            request,
            f"Barang jadi '{fg.name}' tidak dapat dihapus karena sudah memiliki riwayat transaksi mutasi atau pemakaian di project."
        )
        return redirect("finished_good_list")

    name = fg.name
    sku = fg.sku
    fg.delete()
    messages.success(request, f"Barang jadi '{name}' ({sku}) berhasil dihapus.")
    return redirect("finished_good_list")


def finished_good_export_excel(request):
    """
    Export Seluruh Katalog Master Barang Jadi ke Format Excel (.xlsx)
    """
    query = request.GET.get("q", "").strip()
    category = request.GET.get("category", "").strip()
    low_stock = request.GET.get("low_stock", "").strip()

    fg_qs = FinishedGood.objects.all().order_by("name")
    if query:
        fg_qs = fg_qs.filter(
            Q(name__icontains=query) |
            Q(sku__icontains=query) |
            Q(category__icontains=query) |
            Q(notes__icontains=query)
        )
    if category:
        fg_qs = fg_qs.filter(category=category)
    if low_stock == "1":
        fg_qs = fg_qs.filter(current_stock__lte=F("minimum_stock"))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Master Barang Jadi"

    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True)

    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")  # Indigo 600
    total_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")   # Indigo 50
    alert_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")   # Rose 100

    thin_border_side = Side(style="thin", color="CBD5E1")
    border_all = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # Header Title
    ws.merge_cells("A1:K1")
    ws["A1"] = "KATALOG PERSEDIAAN MASTER BARANG JADI (FINISHED GOODS)"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 25

    generated_date = timezone.localtime(timezone.now()).strftime("%d %B %Y, %H:%M")
    ws["A2"] = f"Diekspor pada: {generated_date} | Total Produk: {fg_qs.count()} SKU"
    ws["A2"].font = subtitle_font
    ws.row_dimensions[2].height = 18

    headers = [
        "No", "SKU / Kode", "Nama Barang Jadi", "Kategori", "Satuan",
        "Standard Cost (Rp)", "Stok Fisik", "Batas Minimum", "Status Stok",
        "Total Valuasi (Rp)", "Catatan"
    ]
    ws.append([])
    ws.append(headers)
    header_row_idx = ws.max_row
    ws.row_dimensions[header_row_idx].height = 22

    for col_num in range(1, len(headers) + 1):
        c = ws.cell(row=header_row_idx, column=col_num)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_all

    total_units_sum = Decimal(0)
    total_val_sum = Decimal(0)

    for idx, fg in enumerate(fg_qs, start=1):
        val = fg.total_inventory_value
        total_units_sum += fg.current_stock
        total_val_sum += val
        is_low = fg.current_stock <= fg.minimum_stock
        status_str = "KRITIS / RENDAH" if is_low else "AMAN"

        row = [
            idx,
            fg.sku,
            fg.name,
            fg.category or "Umum",
            fg.unit,
            float(fg.standard_cost),
            float(fg.current_stock),
            float(fg.minimum_stock),
            status_str,
            float(val),
            fg.notes or "-"
        ]
        ws.append(row)
        curr_row = ws.max_row
        for c_idx in range(1, len(row) + 1):
            cell = ws.cell(row=curr_row, column=c_idx)
            cell.border = border_all
            if c_idx in [1, 2, 4, 5, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if c_idx == 9 and is_low:
                    cell.fill = alert_fill
                    cell.font = bold_font
            elif c_idx in [6, 7, 8, 10]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if c_idx in [6, 10]:
                    cell.number_format = "#,##0.00"
                else:
                    cell.number_format = "#,##0.00"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Baris Total
    ws.append([])
    summary_row = [
        "TOTAL", "", "", "", "", "",
        float(total_units_sum), "", "",
        float(total_val_sum), ""
    ]
    ws.append(summary_row)
    sum_row_idx = ws.max_row
    ws.merge_cells(start_row=sum_row_idx, start_column=1, end_row=sum_row_idx, end_column=6)
    for c_idx in range(1, len(summary_row) + 1):
        cell = ws.cell(row=sum_row_idx, column=c_idx)
        cell.font = bold_font
        cell.fill = total_fill
        cell.border = border_all
        if c_idx in [7, 10]:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = "#,##0.00"

    col_widths = {
        "A": 6, "B": 14, "C": 30, "D": 16, "E": 12,
        "F": 18, "G": 14, "H": 14, "I": 18, "J": 22, "K": 26
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"Katalog_Barang_Jadi_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def _reconcile_finished_good_stock(fg):
    """
    Menghitung ulang saldo berjalan (balance_after) seluruh riwayat mutasi barang jadi
    secara kronologis dan menyelaraskan nilai fg.current_stock.
    """
    mutations = fg.mutations.all().order_by("created_at", "id")
    running_balance = Decimal(0)
    for m in mutations:
        if m.mutation_type == "IN":
            running_balance += m.quantity
        else:
            running_balance -= m.quantity
        if m.balance_after != running_balance:
            m.balance_after = running_balance
            m.save(update_fields=["balance_after"])
    if fg.current_stock != running_balance:
        fg.current_stock = running_balance
        fg.save(update_fields=["current_stock"])
    return running_balance


def finished_good_mutation_list(request):
    """
    Halaman Menu Mutasi Barang Jadi (Pencatatan Barang Masuk, Rusak, Penyesuaian)
    Lengkap dengan filter periode tanggal, autocomplete combobox produk, 4 kartu KPI,
    dan pagination.
    """
    fg_uuid = request.GET.get("fg", "").strip()
    mutation_type = request.GET.get("type", "").strip()
    query = request.GET.get("q", "").strip()
    start_date_str = request.GET.get("start_date", "").strip()
    end_date_str = request.GET.get("end_date", "").strip()

    mutations_qs = FinishedGoodStockMutation.objects.all().select_related("finished_good", "project")

    if fg_uuid:
        mutations_qs = mutations_qs.filter(finished_good__uuid=fg_uuid)
    if mutation_type:
        mutations_qs = mutations_qs.filter(mutation_type=mutation_type)
    if query:
        mutations_qs = mutations_qs.filter(
            Q(reference_no__icontains=query) |
            Q(notes__icontains=query) |
            Q(finished_good__name__icontains=query) |
            Q(finished_good__sku__icontains=query) |
            Q(project__code__icontains=query) |
            Q(project__name__icontains=query)
        )

    # Filter rentang tanggal
    if start_date_str:
        try:
            start_date = datetime.strptime(start_date_str, "%Y-%m-%d").date()
            start_datetime = timezone.make_aware(datetime.combine(start_date, time.min))
            mutations_qs = mutations_qs.filter(created_at__gte=start_datetime)
        except ValueError:
            start_date_str = ""

    if end_date_str:
        try:
            end_date = datetime.strptime(end_date_str, "%Y-%m-%d").date()
            end_datetime = timezone.make_aware(datetime.combine(end_date, time.max))
            mutations_qs = mutations_qs.filter(created_at__lte=end_datetime)
        except ValueError:
            end_date_str = ""

    # Anotasi nilai transaksi per baris (quantity * standard_cost)
    trans_val_expr = ExpressionWrapper(
        F("quantity") * F("finished_good__standard_cost"),
        output_field=DecimalField(max_digits=15, decimal_places=2)
    )
    mutations_qs = mutations_qs.annotate(transaction_value=trans_val_expr)

    # Statistik ringkasan 4 KPI
    stats = mutations_qs.aggregate(
        total_in=Count("id", filter=Q(mutation_type="IN")),
        total_out=Count("id", filter=Q(mutation_type="OUT")),
        total_in_qty=Sum("quantity", filter=Q(mutation_type="IN")),
        total_out_qty=Sum("quantity", filter=Q(mutation_type="OUT")),
        total_val=Sum("transaction_value")
    )
    total_in_count = stats["total_in"] or 0
    total_out_count = stats["total_out"] or 0
    total_in_qty = stats["total_in_qty"] or Decimal(0)
    total_out_qty = stats["total_out_qty"] or Decimal(0)
    total_val = stats["total_val"] or Decimal(0)
    total_count = total_in_count + total_out_count

    finished_goods = FinishedGood.objects.all().order_by("name")

    selected_fg_obj = None
    if fg_uuid:
        selected_fg_obj = FinishedGood.objects.filter(uuid=fg_uuid).first()

    finished_goods_json = [
        {
            "uuid": str(fg.uuid),
            "sku": fg.sku,
            "name": fg.name,
            "unit": fg.unit,
            "current_stock": float(fg.current_stock),
            "standard_cost": float(fg.standard_cost),
        }
        for fg in finished_goods
    ]

    # Pagination 20 data per halaman
    paginator = Paginator(mutations_qs, 20)
    page_number = request.GET.get("page", 1)
    mutations = paginator.get_page(page_number)

    return render(request, "hpp/stock_mutation_list.html", {
        "mutations": mutations,
        "finished_goods": finished_goods,
        "finished_goods_json": json.dumps(finished_goods_json),
        "selected_fg": fg_uuid,
        "selected_fg_obj": selected_fg_obj,
        "selected_type": mutation_type,
        "query": query,
        "start_date_str": start_date_str,
        "end_date_str": end_date_str,
        "total_count": total_count,
        "total_in_count": total_in_count,
        "total_out_count": total_out_count,
        "total_in_qty": total_in_qty,
        "total_out_qty": total_out_qty,
        "total_val": total_val,
    })


@transaction.atomic
def finished_good_mutation_create(request):
    """
    Pencatatan transaksi mutasi barang masuk, barang rusak, atau penyesuaian opname.
    """
    if request.method == "POST":
        fg_uuid = request.POST.get("finished_good_uuid", "").strip()
        mutation_type = request.POST.get("mutation_type", "IN").strip()
        reason_type = request.POST.get("reason_type", "masuk").strip()
        quantity_str = request.POST.get("quantity", "0").strip()
        reference_no = request.POST.get("reference_no", "").strip()
        notes = request.POST.get("notes", "").strip()

        try:
            quantity = Decimal(quantity_str)
        except Exception:
            quantity = Decimal(0)

        fg = get_object_or_404(FinishedGood, uuid=fg_uuid)

        if quantity <= 0:
            messages.error(request, "Jumlah mutasi harus lebih besar dari 0.")
            return redirect("stock_mutation_list")

        if mutation_type == "IN":
            fg.current_stock += quantity
            prefix = "[MASUK]"
        else:
            if fg.current_stock < quantity:
                messages.error(
                    request,
                    f"Stok {fg.name} tidak mencukupi (Tersedia: {fg.current_stock} {fg.unit}, diminta: {quantity} {fg.unit})."
                )
                return redirect("stock_mutation_list")
            fg.current_stock -= quantity
            prefix = "[RUSAK/OUT]" if reason_type == "rusak" else "[KOREKSI/OUT]"

        fg.save(update_fields=["current_stock"])

        full_notes = f"{prefix} {notes}".strip()
        FinishedGoodStockMutation.objects.create(
            finished_good=fg,
            mutation_type=mutation_type,
            quantity=quantity,
            balance_after=fg.current_stock,
            reference_no=reference_no or f"MUT-{fg.sku}",
            notes=full_notes
        )

        messages.success(
            request,
            f"Mutasi stok {fg.name} ({mutation_type} {quantity} {fg.unit}) berhasil dicatat. Saldo terkini: {fg.current_stock} {fg.unit}."
        )
    return redirect("stock_mutation_list")


def finished_good_mutation_export_excel(request):
    """
    Export Rekapitulasi Log Mutasi Stok Barang Jadi ke File Excel (.xlsx)
    Sesuai filter pencarian, produk barang jadi, tipe mutasi, dan rentang tanggal.
    """
    fg_uuid = request.GET.get("fg", "").strip()
    mutation_type = request.GET.get("type", "").strip()
    query = request.GET.get("q", "").strip()
    start_date_str = request.GET.get("start_date", "").strip()
    end_date_str = request.GET.get("end_date", "").strip()

    mutations_qs = FinishedGoodStockMutation.objects.all().select_related("finished_good", "project")

    selected_fg = None
    if fg_uuid:
        selected_fg = FinishedGood.objects.filter(uuid=fg_uuid).first()
        if selected_fg:
            mutations_qs = mutations_qs.filter(finished_good=selected_fg)

    if mutation_type:
        mutations_qs = mutations_qs.filter(mutation_type=mutation_type)

    if query:
        mutations_qs = mutations_qs.filter(
            Q(reference_no__icontains=query) |
            Q(notes__icontains=query) |
            Q(finished_good__name__icontains=query) |
            Q(finished_good__sku__icontains=query) |
            Q(project__code__icontains=query) |
            Q(project__name__icontains=query)
        )

    if start_date_str:
        try:
            start_date = datetime.strptime(start_date_str, "%Y-%m-%d").date()
            start_datetime = timezone.make_aware(datetime.combine(start_date, time.min))
            mutations_qs = mutations_qs.filter(created_at__gte=start_datetime)
        except ValueError:
            start_date_str = ""

    if end_date_str:
        try:
            end_date = datetime.strptime(end_date_str, "%Y-%m-%d").date()
            end_datetime = timezone.make_aware(datetime.combine(end_date, time.max))
            mutations_qs = mutations_qs.filter(created_at__lte=end_datetime)
        except ValueError:
            end_date_str = ""

    mutations = mutations_qs.order_by("-created_at")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mutasi Barang Jadi"

    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True)

    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")  # Indigo 600
    in_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")       # Emerald 100
    out_fill = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")      # Rose 100
    total_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")    # Indigo 50

    thin_border_side = Side(style="thin", color="CBD5E1")
    border_all = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # Title
    ws.merge_cells("A1:L1")
    ws["A1"] = "LOG MUTASI PERSEDIAAN BARANG JADI (FINISHED GOODS)"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 25

    # Subtitle / Filter info
    fg_label = f"[{selected_fg.sku}] {selected_fg.name}" if selected_fg else "Semua Barang Jadi"
    type_label = "Semua Tipe" if not mutation_type else ("Barang Masuk (IN)" if mutation_type == "IN" else "Barang Keluar (OUT)")
    periode_label = f"{start_date_str or 'Awal'} s/d {end_date_str or 'Sekarang'}"
    generated_date = timezone.localtime(timezone.now()).strftime("%d %B %Y, %H:%M")

    ws["A2"] = f"Filter Produk: {fg_label} | Tipe: {type_label} | Periode: {periode_label} | Unduh: {generated_date}"
    ws["A2"].font = subtitle_font
    ws.row_dimensions[2].height = 18

    headers = [
        "No", "Waktu Transaksi", "No. Referensi", "SKU", "Nama Barang Jadi",
        "Arah Mutasi", "Jumlah Mutasi", "Satuan", "Saldo Akhir",
        "Standard Cost (Rp)", "Nilai Transaksi (Rp)", "Keterangan / Project"
    ]
    ws.append([])
    ws.append(headers)
    header_row_idx = ws.max_row
    ws.row_dimensions[header_row_idx].height = 22

    for col_num in range(1, len(headers) + 1):
        c = ws.cell(row=header_row_idx, column=col_num)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_all

    total_in_qty = Decimal(0)
    total_out_qty = Decimal(0)
    total_value_sum = Decimal(0)

    for idx, m in enumerate(mutations, start=1):
        local_time = timezone.localtime(m.created_at).strftime("%d/%m/%Y %H:%M")
        cost = m.finished_good.standard_cost
        trans_val = m.quantity * cost
        total_value_sum += trans_val

        if m.mutation_type == "IN":
            total_in_qty += m.quantity
            type_display = "MASUK (IN)"
            fill_type = in_fill
        else:
            total_out_qty += m.quantity
            type_display = "KELUAR (OUT)"
            fill_type = out_fill

        ref_str = m.reference_no or "-"
        if m.project:
            ref_str = f"{ref_str} ({m.project.code})"

        row = [
            idx,
            local_time,
            ref_str,
            m.finished_good.sku,
            m.finished_good.name,
            type_display,
            float(m.quantity),
            m.finished_good.unit,
            float(m.balance_after),
            float(cost),
            float(trans_val),
            m.notes or "-"
        ]
        ws.append(row)
        curr_row = ws.max_row
        for c_idx in range(1, len(row) + 1):
            cell = ws.cell(row=curr_row, column=c_idx)
            cell.border = border_all
            if c_idx in [1, 2, 4, 6, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if c_idx == 6:
                    cell.fill = fill_type
                    cell.font = bold_font
            elif c_idx in [7, 9, 10, 11]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.00"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Baris Summary Total
    ws.append([])
    summary_row = [
        "TOTAL PERIODE INI", "", "", "", "", "",
        "", "", "", "",
        float(total_value_sum), ""
    ]
    ws.append(summary_row)
    sum_row_idx = ws.max_row
    ws.merge_cells(start_row=sum_row_idx, start_column=1, end_row=sum_row_idx, end_column=10)
    for c_idx in range(1, len(summary_row) + 1):
        cell = ws.cell(row=sum_row_idx, column=c_idx)
        cell.font = bold_font
        cell.fill = total_fill
        cell.border = border_all
        if c_idx == 11:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = "#,##0.00"

    col_widths = {
        "A": 6, "B": 18, "C": 20, "D": 14, "E": 28, "F": 16,
        "G": 14, "H": 10, "I": 14, "J": 18, "K": 20, "L": 35
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"Mutasi_Barang_Jadi_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def _get_finished_good_stock_card_data(fg_uuid, start_date_str, end_date_str):
    """
    Helper untuk menghitung data buku besar kartu stok barang jadi:
    - selected_fg
    - beginning_balance
    - total_in (periode)
    - total_out (periode)
    - ending_balance (periode)
    - ending_valuation (periode)
    - mutations_qs
    - start_date, end_date
    """
    finished_goods = FinishedGood.objects.all().order_by("name")
    if fg_uuid:
        selected_fg = FinishedGood.objects.filter(uuid=fg_uuid).first()
    else:
        selected_fg = finished_goods.first()

    start_date = None
    end_date = None
    start_datetime = None
    end_datetime = None

    if start_date_str:
        try:
            start_date = datetime.strptime(start_date_str, "%Y-%m-%d").date()
            start_datetime = timezone.make_aware(datetime.combine(start_date, time.min))
        except ValueError:
            start_date_str = ""

    if end_date_str:
        try:
            end_date = datetime.strptime(end_date_str, "%Y-%m-%d").date()
            end_datetime = timezone.make_aware(datetime.combine(end_date, time.max))
        except ValueError:
            end_date_str = ""

    beginning_balance = Decimal(0)
    total_in = Decimal(0)
    total_out = Decimal(0)
    ending_balance = Decimal(0)
    ending_valuation = Decimal(0)
    mutations_qs = FinishedGoodStockMutation.objects.none()

    if selected_fg:
        # Hitung Saldo Awal sebelum rentang tanggal yang dipilih
        if start_datetime:
            prior_mutations = selected_fg.mutations.filter(created_at__lt=start_datetime)
            last_prior = prior_mutations.order_by("-created_at", "-id").first()
            if last_prior:
                beginning_balance = last_prior.balance_after
            else:
                beginning_balance = Decimal(0)
        else:
            beginning_balance = Decimal(0)

        # Mutasi dalam periode
        mutations_qs = selected_fg.mutations.all().select_related("project").order_by("created_at", "id")
        if start_datetime:
            mutations_qs = mutations_qs.filter(created_at__gte=start_datetime)
        if end_datetime:
            mutations_qs = mutations_qs.filter(created_at__lte=end_datetime)

        totals = mutations_qs.aggregate(
            total_in=Sum("quantity", filter=Q(mutation_type="IN")),
            total_out=Sum("quantity", filter=Q(mutation_type="OUT"))
        )
        total_in = totals["total_in"] or Decimal(0)
        total_out = totals["total_out"] or Decimal(0)

        if start_datetime:
            ending_balance = beginning_balance + total_in - total_out
        else:
            ending_balance = selected_fg.current_stock

        ending_valuation = ending_balance * selected_fg.standard_cost

    return {
        "finished_goods": finished_goods,
        "selected_fg": selected_fg,
        "start_date_str": start_date_str,
        "end_date_str": end_date_str,
        "start_date": start_date,
        "end_date": end_date,
        "beginning_balance": beginning_balance,
        "total_in": total_in,
        "total_out": total_out,
        "ending_balance": ending_balance,
        "ending_valuation": ending_valuation,
        "mutations_qs": mutations_qs,
    }


def finished_good_stock_card_index(request):
    """
    Menu Kartu Stok Barang Jadi Standalone (Pencarian Combobox, Filter Tanggal, Saldo Awal, Print, Export)
    """
    fg_uuid = request.GET.get("fg", "").strip()
    start_date_str = request.GET.get("start_date", "").strip()
    end_date_str = request.GET.get("end_date", "").strip()
    page_number = request.GET.get("page", 1)

    data = _get_finished_good_stock_card_data(fg_uuid, start_date_str, end_date_str)
    selected_fg = data["selected_fg"]
    mutations_qs = data["mutations_qs"]

    # Pagination 25 data per halaman
    paginator = Paginator(mutations_qs, 25)
    mutations = paginator.get_page(page_number)

    # Anotasi dinamis transaction_value untuk setiap mutasi di halaman
    for m in mutations:
        m.transaction_value = m.quantity * (selected_fg.standard_cost if selected_fg else Decimal(0))

    finished_goods_json = [
        {
            "uuid": str(fg.uuid),
            "sku": fg.sku,
            "name": fg.name,
            "unit": fg.unit,
            "current_stock": float(fg.current_stock),
            "standard_cost": float(fg.standard_cost),
        }
        for fg in data["finished_goods"]
    ]

    return render(request, "hpp/stock_card_index.html", {
        "finished_goods": data["finished_goods"],
        "finished_goods_json": json.dumps(finished_goods_json),
        "selected_fg": selected_fg,
        "mutations": mutations,
        "start_date": data["start_date"],
        "end_date": data["end_date"],
        "start_date_str": data["start_date_str"],
        "end_date_str": data["end_date_str"],
        "beginning_balance": data["beginning_balance"],
        "total_in": data["total_in"],
        "total_out": data["total_out"],
        "ending_balance": data["ending_balance"],
        "ending_valuation": data["ending_valuation"],
    })


def finished_good_stock_card(request, uuid):
    """
    Kartu Stok Barang Jadi Tunggal -> Redirect ke index dengan parameter ?fg=uuid
    agar selalu menikmati fitur lengkap (filter tanggal, export excel, print view, combobox)
    """
    return redirect(f"/stock/cards/?fg={uuid}")


def finished_good_stock_card_export_excel(request):
    """
    Export Kartu Stok Barang Jadi ke format Excel (.xlsx) menggunakan openpyxl
    """
    fg_uuid = request.GET.get("fg", "").strip()
    start_date_str = request.GET.get("start_date", "").strip()
    end_date_str = request.GET.get("end_date", "").strip()

    data = _get_finished_good_stock_card_data(fg_uuid, start_date_str, end_date_str)
    selected_fg = data["selected_fg"]
    if not selected_fg:
        messages.error(request, "Barang jadi tidak ditemukan.")
        return redirect("stock_card_index")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Kartu Stok"

    # Fonts & Styles
    title_font = Font(name="Calibri", size=15, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    bold_font = Font(name="Calibri", size=10, bold=True)
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")  # Indigo 600
    amber_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")   # Amber 100
    in_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")      # Emerald 100
    out_fill = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")     # Rose 100
    total_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")   # Indigo 50

    thin_border = Side(style="thin", color="CBD5E1")
    border_all = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    # Title
    ws.merge_cells("A1:I1")
    ws["A1"] = "KARTU STOK PERSEDIAAN BARANG JADI (FINISHED GOODS)"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 25

    generated_date = timezone.localtime(timezone.now()).strftime("%d %B %Y, %H:%M")
    periode_label = f"{data['start_date_str'] or 'Awal'} s/d {data['end_date_str'] or 'Sekarang'}"

    # Metadata Produk & Periode
    ws["A2"] = f"Produk: [{selected_fg.sku}] {selected_fg.name} | Satuan: {selected_fg.unit} | Cost Standar: Rp {selected_fg.standard_cost:,.2f}"
    ws["A2"].font = bold_font
    ws["A3"] = f"Periode: {periode_label} | Saldo Awal: {data['beginning_balance']:,.2f} {selected_fg.unit} | Saldo Akhir: {data['ending_balance']:,.2f} {selected_fg.unit} | Valuasi: Rp {data['ending_valuation']:,.2f} | Dicetak: {generated_date}"
    ws["A3"].font = subtitle_font

    headers = [
        "No", "Waktu Transaksi", "No. Referensi / Project", "Keterangan",
        f"Masuk ({selected_fg.unit})", f"Keluar ({selected_fg.unit})",
        f"Saldo ({selected_fg.unit})", "Cost Standar (Rp)", "Nilai Transaksi (Rp)"
    ]
    ws.append([])
    ws.append(headers)
    header_row_idx = ws.max_row
    ws.row_dimensions[header_row_idx].height = 24

    for col_num in range(1, len(headers) + 1):
        c = ws.cell(row=header_row_idx, column=col_num)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border_all

    # Baris Saldo Awal jika ada filter tanggal
    if data["start_date"]:
        row_init = [
            "-",
            f"{data['start_date_str']} 00:00",
            "SALDO AWAL",
            f"Saldo akumulasi persediaan sebelum {data['start_date_str']}",
            "",
            "",
            float(data["beginning_balance"]),
            float(selected_fg.standard_cost),
            float(data["beginning_balance"] * selected_fg.standard_cost)
        ]
        ws.append(row_init)
        curr_row = ws.max_row
        for c_idx in range(1, len(row_init) + 1):
            cell = ws.cell(row=curr_row, column=c_idx)
            cell.border = border_all
            cell.fill = amber_fill
            cell.font = bold_font
            if c_idx in [1, 2, 3]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in [5, 6, 7, 8, 9]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.00"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    total_in_sum = Decimal(0)
    total_out_sum = Decimal(0)
    total_val_sum = Decimal(0)

    for idx, m in enumerate(data["mutations_qs"], start=1):
        local_time = timezone.localtime(m.created_at).strftime("%d/%m/%Y %H:%M")
        cost = selected_fg.standard_cost
        trans_val = m.quantity * cost
        total_val_sum += trans_val

        in_qty = float(m.quantity) if m.mutation_type == "IN" else ""
        out_qty = float(m.quantity) if m.mutation_type == "OUT" else ""

        if m.mutation_type == "IN":
            total_in_sum += m.quantity
        else:
            total_out_sum += m.quantity

        ref_str = m.reference_no or "-"
        if m.project:
            ref_str = f"{ref_str} ({m.project.code})"

        row = [
            idx,
            local_time,
            ref_str,
            m.notes or "-",
            in_qty,
            out_qty,
            float(m.balance_after),
            float(cost),
            float(trans_val)
        ]
        ws.append(row)
        curr_row = ws.max_row
        for c_idx in range(1, len(row) + 1):
            cell = ws.cell(row=curr_row, column=c_idx)
            cell.border = border_all
            if c_idx in [1, 2]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx in [5, 6, 7, 8, 9]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.00"
                if c_idx == 5 and in_qty:
                    cell.fill = in_fill
                    cell.font = bold_font
                elif c_idx == 6 and out_qty:
                    cell.fill = out_fill
                    cell.font = bold_font
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Baris Summary Total
    ws.append([])
    summary_row = [
        "TOTAL PERIODE INI", "", "", "",
        float(total_in_sum),
        float(total_out_sum),
        float(data["ending_balance"]),
        "",
        float(total_val_sum)
    ]
    ws.append(summary_row)
    sum_row_idx = ws.max_row
    ws.merge_cells(start_row=sum_row_idx, start_column=1, end_row=sum_row_idx, end_column=4)
    for c_idx in range(1, len(summary_row) + 1):
        cell = ws.cell(row=sum_row_idx, column=c_idx)
        cell.font = bold_font
        cell.fill = total_fill
        cell.border = border_all
        if c_idx in [5, 6, 7, 9]:
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = "#,##0.00"

    col_widths = {
        "A": 6, "B": 18, "C": 22, "D": 32, "E": 15,
        "F": 15, "G": 16, "H": 18, "I": 20
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    filename = f"Kartu_Stok_{selected_fg.sku}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
