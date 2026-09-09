from django.shortcuts import render, get_object_or_404, redirect
from django.contrib import messages
from django.db import transaction
from django.db.models import Q
from django.core.paginator import Paginator
from django.http import HttpResponse
from datetime import datetime
from decimal import Decimal
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import (
    UnitMaster,
    Project, BOMItem, ProjectLabor, ProjectOverhead, FinishedGood,
    ProjectFinishedGood, BOMItemRealization, LaborRealization,
    FinishedGoodRealization, OverheadRealization, RawMaterial
)
from .views_auth import menu_permission_required


def _clean_decimal(val_str, default=Decimal(0)):
    if not val_str:
        return default
    if isinstance(val_str, (int, float, Decimal)):
        return Decimal(str(val_str))
    cleaned = str(val_str).strip()
    cleaned = cleaned.replace("Rp", "").replace("rp", "").replace("RP", "").strip()
    cleaned = cleaned.replace(" ", "")
    if "." in cleaned and "," in cleaned:
        cleaned = cleaned.replace(".", "").replace(",", ".")
    elif "," in cleaned and "." not in cleaned:
        cleaned = cleaned.replace(",", ".")
    elif "." in cleaned:
        parts = cleaned.split(".")
        if len(parts) > 2:
            cleaned = "".join(parts[:-1]) + "." + parts[-1]
        elif len(parts) == 2 and len(parts[1]) == 3:
            cleaned = cleaned.replace(".", "")
    try:
        return Decimal(cleaned)
    except Exception:
        return default


@menu_permission_required("MENU_PROJECT_REALIZATION")
def project_realization_list(request):
    query = request.GET.get("q", "").strip()
    status = request.GET.get("status", "").strip()
    budget_status = request.GET.get("budget_status", "all").strip()
    page_num = request.GET.get("page", 1)

    projects_qs = Project.objects.all().order_by("-created_at")
    if query:
        projects_qs = projects_qs.filter(Q(name__icontains=query) | Q(code__icontains=query) | Q(customer_name__icontains=query))
    if status:
        projects_qs = projects_qs.filter(status=status)

    all_projects = list(projects_qs)
    total_projects_count = len(all_projects)
    total_burned_cost = Decimal(0)
    total_budget_plan = Decimal(0)
    total_contract_value = Decimal(0)
    over_budget_count = 0
    on_budget_count = 0
    no_realization_count = 0

    enriched_projects = []
    for p in all_projects:
        est_hpp = p.total_hpp_estimated
        act_hpp = p.total_hpp_actual
        contract_val = p.contract_value
        diff = act_hpp - est_hpp
        remaining = est_hpp - act_hpp

        total_burned_cost += act_hpp
        total_budget_plan += est_hpp
        total_contract_value += contract_val

        if est_hpp > 0:
            burn_pct = (act_hpp / est_hpp) * Decimal(100)
        else:
            burn_pct = Decimal(0) if act_hpp == 0 else Decimal(100)

        if act_hpp == 0:
            b_status = "no_realization"
            no_realization_count += 1
        elif act_hpp > est_hpp:
            b_status = "over_budget"
            over_budget_count += 1
        else:
            b_status = "on_budget"
            on_budget_count += 1

        p.calc_est_hpp = est_hpp
        p.calc_act_hpp = act_hpp
        p.calc_diff = diff
        p.calc_remaining = remaining
        p.calc_burn_pct = round(burn_pct, 1)
        p.calc_budget_status = b_status
        enriched_projects.append(p)

    # Filter berdasarkan status anggaran
    if budget_status == "over_budget":
        filtered_projects = [p for p in enriched_projects if p.calc_budget_status == "over_budget"]
    elif budget_status == "on_budget":
        filtered_projects = [p for p in enriched_projects if p.calc_budget_status == "on_budget"]
    elif budget_status == "no_realization":
        filtered_projects = [p for p in enriched_projects if p.calc_budget_status == "no_realization"]
    else:
        filtered_projects = enriched_projects

    overall_act_margin = Decimal(0)
    if total_contract_value > 0:
        overall_act_margin = ((total_contract_value - total_burned_cost) / total_contract_value) * Decimal(100)

    # Paginasi 20 item per halaman
    paginator = Paginator(filtered_projects, 20)
    page_obj = paginator.get_page(page_num)

    return render(request, "hpp/project_realization_list.html", {
        "projects": page_obj,
        "query": query,
        "selected_status": status,
        "selected_budget_status": budget_status,
        "status_choices": Project.STATUS_CHOICES,
        "total_projects_count": total_projects_count,
        "filtered_count": len(filtered_projects),
        "total_burned_cost": total_burned_cost,
        "total_budget_plan": total_budget_plan,
        "over_budget_count": over_budget_count,
        "on_budget_count": on_budget_count,
        "no_realization_count": no_realization_count,
        "overall_act_margin": overall_act_margin,
    })


def project_realization_export_excel(request):
    query = request.GET.get("q", "").strip()
    status = request.GET.get("status", "").strip()
    budget_status = request.GET.get("budget_status", "all").strip()

    projects_qs = Project.objects.all().order_by("-created_at")
    if query:
        projects_qs = projects_qs.filter(Q(name__icontains=query) | Q(code__icontains=query) | Q(customer_name__icontains=query))
    if status:
        projects_qs = projects_qs.filter(status=status)

    enriched = []
    for p in projects_qs:
        est_hpp = p.total_hpp_estimated
        act_hpp = p.total_hpp_actual
        diff = act_hpp - est_hpp
        remaining = est_hpp - act_hpp
        burn_pct = (act_hpp / est_hpp * Decimal(100)) if est_hpp > 0 else (Decimal(0) if act_hpp == 0 else Decimal(100))
        
        if act_hpp == 0:
            b_status = "no_realization"
            status_label = "Belum Ada Realisasi"
        elif act_hpp > est_hpp:
            b_status = "over_budget"
            status_label = "OVER BUDGET"
        else:
            b_status = "on_budget"
            status_label = "ON BUDGET / HEMAT"

        p.calc_est_hpp = est_hpp
        p.calc_act_hpp = act_hpp
        p.calc_diff = diff
        p.calc_remaining = remaining
        p.calc_burn_pct = burn_pct
        p.calc_budget_status = b_status
        p.calc_status_label = status_label

        if budget_status == "over_budget" and b_status != "over_budget":
            continue
        if budget_status == "on_budget" and b_status != "on_budget":
            continue
        if budget_status == "no_realization" and b_status != "no_realization":
            continue

        enriched.append(p)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Rekap Realisasi HPP"
    ws.views.sheetView[0].showGridLines = True

    # Styling
    font_family = "Calibri"
    thin_side = Side(border_style="thin", color="CBD5E1")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    double_bottom_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=Side(border_style="double", color="0F172A"))

    title_font = Font(name=font_family, size=14, bold=True, color="1E1B4B")
    tbl_header_font = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    tbl_header_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
    
    total_font = Font(name=font_family, size=10, bold=True, color="0F172A")
    total_fill = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")

    currency_fmt = '"Rp" #,##0;("Rp" #,##0);"-"'
    percent_fmt = '0.00%'

    # Header title
    ws.append(["REKAPITULASI PENYERAPAN ANGGARAN & REALISASI HPP PROJECT"])
    ws["A1"].font = title_font
    ws.append([f"Tanggal Ekspor: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Total Data: {len(enriched)} Project"])
    ws.append([])

    # Table Header
    headers = [
        "No", "Kode Project", "Nama Project", "Customer", "Status Workflow",
        "Nilai Kontrak", "Anggaran (Est. HPP)", "Realisasi (Aktual HPP)",
        "Sisa Anggaran", "Deviasi Biaya", "Penyerapan (%)",
        "Status Anggaran", "Gross Profit (Aktual)", "Margin Aktual (%)"
    ]
    ws.append(headers)
    hdr_row = ws.max_row
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=hdr_row, column=c)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c in [1, 2, 5, 11, 12] else ("right" if c >= 6 else "left"))

    # Rows
    tot_contract = Decimal(0)
    tot_est = Decimal(0)
    tot_act = Decimal(0)
    tot_remaining = Decimal(0)
    tot_diff = Decimal(0)
    tot_gp = Decimal(0)

    for idx, p in enumerate(enriched, start=1):
        tot_contract += p.contract_value
        tot_est += p.calc_est_hpp
        tot_act += p.calc_act_hpp
        tot_remaining += p.calc_remaining
        tot_diff += p.calc_diff
        tot_gp += p.act_gross_profit

        burn_ratio = float(p.calc_burn_pct) / 100.0
        margin_ratio = float(p.act_margin_pct) / 100.0

        ws.append([
            idx,
            p.code,
            p.name,
            p.customer_name or "-",
            p.get_status_display(),
            float(p.contract_value),
            float(p.calc_est_hpp),
            float(p.calc_act_hpp),
            float(p.calc_remaining),
            float(p.calc_diff),
            burn_ratio,
            p.calc_status_label,
            float(p.act_gross_profit),
            margin_ratio
        ])
        curr_row = ws.max_row
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=curr_row, column=c)
            cell.border = thin_border
            if c in [6, 7, 8, 9, 10, 13]:
                cell.number_format = currency_fmt
                cell.alignment = Alignment(horizontal="right")
            elif c in [11, 14]:
                cell.number_format = percent_fmt
                cell.alignment = Alignment(horizontal="right")
            elif c in [1, 2, 5, 12]:
                cell.alignment = Alignment(horizontal="center")
            
            # Status colors
            if c == 12:
                if p.calc_budget_status == "over_budget":
                    cell.font = Font(name=font_family, size=9, bold=True, color="B91C1C")
                    cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                elif p.calc_budget_status == "on_budget":
                    cell.font = Font(name=font_family, size=9, bold=True, color="047857")
                    cell.fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")

    # Grand Total Row
    tot_burn_ratio = float(tot_act / tot_est) if tot_est > 0 else 0.0
    tot_margin_ratio = float(tot_gp / tot_contract) if tot_contract > 0 else 0.0
    ws.append([
        "TOTAL", "", "", "", "",
        float(tot_contract),
        float(tot_est),
        float(tot_act),
        float(tot_remaining),
        float(tot_diff),
        tot_burn_ratio,
        "",
        float(tot_gp),
        tot_margin_ratio
    ])
    tot_row = ws.max_row
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=tot_row, column=c)
        cell.font = total_font
        cell.fill = total_fill
        cell.border = double_bottom_border
        if c in [6, 7, 8, 9, 10, 13]:
            cell.number_format = currency_fmt
            cell.alignment = Alignment(horizontal="right")
        elif c in [11, 14]:
            cell.number_format = percent_fmt
            cell.alignment = Alignment(horizontal="right")
        elif c == 1:
            cell.alignment = Alignment(horizontal="center")

    # Autofit
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.value is not None:
                val_str = str(cell.value)
                if "\n" in val_str:
                    val_str = max(val_str.split("\n"), key=len)
                if len(val_str) > max_len:
                    max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 46)

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    filename = f"Rekap_Realisasi_Project_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


def project_realization_detail(request, uuid):
    project = get_object_or_404(Project, uuid=uuid)
    bom_items = project.bom_items.filter(item_type="material").select_related("material_master", "raw_material")
    labor_items = project.labor_items.all().select_related("labor_master")
    fg_items = project.finished_good_items.all().select_related("finished_good")
    overhead_items = project.overhead_items.all()

    master_finished_goods = FinishedGood.objects.all().order_by("name")
    master_raw_materials = RawMaterial.objects.all().prefetch_related("conversions").order_by("name")

    bom_realizations = project.bom_realizations.all().select_related("bom_item", "raw_material")
    labor_realizations = project.labor_realizations.all().select_related("project_labor")
    fg_realizations = project.finished_good_realizations.all().select_related("project_fg", "finished_good")
    overhead_realizations = project.overhead_realizations.all().select_related("project_overhead")

    # Serialize Master Raw Materials
    raw_materials_list = []
    for rm in master_raw_materials:
        raw_materials_list.append({
            "uuid": str(rm.uuid),
            "code": rm.code,
            "name": rm.name,
            "category": rm.category or "Umum",
            "stock_unit": rm.stock_unit,
            "current_stock": float(rm.current_stock),
            "last_purchase_price": float(rm.last_purchase_price),
            "conversions": [
                {
                    "unit_name": c.unit_name,
                    "conversion_factor": float(c.conversion_factor)
                } for c in rm.conversions.all()
            ]
        })

    # Serialize Planned BOM Items
    bom_items_list = []
    for b in bom_items:
        rm = b.raw_material
        if not rm and b.name:
            rm = RawMaterial.objects.filter(name__iexact=b.name.strip()).first()
        rm_uuid = str(rm.uuid) if rm else ""
        current_stock = float(rm.current_stock) if rm else 0
        stock_unit = rm.stock_unit if rm else b.unit
        bom_items_list.append({
            "uuid": str(b.uuid),
            "name": b.name,
            "unit": b.unit,
            "est_qty": float(b.est_qty),
            "est_unit_cost": float(b.est_unit_cost),
            "raw_material_uuid": rm_uuid,
            "current_stock": current_stock,
            "stock_unit": stock_unit,
        })

    # Serialize Planned Labor Items
    labor_items_list = []
    for l in labor_items:
        labor_items_list.append({
            "uuid": str(l.uuid),
            "role_name": l.role_name,
            "unit": l.unit,
            "est_quantity": float(l.est_quantity),
            "est_rate": float(l.est_rate),
        })

    # Serialize Master Finished Goods
    finished_goods_list = []
    for fg in master_finished_goods:
        finished_goods_list.append({
            "uuid": str(fg.uuid),
            "sku": fg.sku,
            "name": fg.name,
            "unit": fg.unit,
            "current_stock": float(fg.current_stock),
            "standard_cost": float(fg.standard_cost),
        })

    # Serialize Planned FG Items
    fg_items_list = []
    for pfg in fg_items:
        fg_items_list.append({
            "uuid": str(pfg.uuid),
            "fg_uuid": str(pfg.finished_good.uuid),
            "sku": pfg.finished_good.sku,
            "name": pfg.finished_good.name,
            "unit": pfg.finished_good.unit,
            "est_qty": float(pfg.est_qty),
            "est_unit_cost": float(pfg.est_unit_cost),
            "current_stock": float(pfg.finished_good.current_stock),
        })

    # Serialize Planned Overhead Items
    overhead_items_list = []
    for o in overhead_items:
        overhead_items_list.append({
            "uuid": str(o.uuid),
            "name": o.name,
            "est_cost": float(o.est_cost),
            "notes": o.notes or "",
        })

    unit_materials = UnitMaster.objects.filter(category="raw_material", is_active=True).order_by("name")
    unit_labors = UnitMaster.objects.filter(category="labor", is_active=True).order_by("name")

    return render(request, "hpp/project_realization_detail.html", {
        "project": project,
        "bom_items": bom_items,
        "labor_items": labor_items,
        "fg_items": fg_items,
        "overhead_items": overhead_items,
        "master_finished_goods": master_finished_goods,
        "master_raw_materials": master_raw_materials,
        "bom_realizations": bom_realizations,
        "labor_realizations": labor_realizations,
        "fg_realizations": fg_realizations,
        "overhead_realizations": overhead_realizations,
        "raw_materials_json": json.dumps(raw_materials_list),
        "bom_items_json": json.dumps(bom_items_list),
        "labor_items_json": json.dumps(labor_items_list),
        "finished_goods_json": json.dumps(finished_goods_list),
        "fg_items_json": json.dumps(fg_items_list),
        "overhead_items_json": json.dumps(overhead_items_list),
        "unit_materials": unit_materials,
        "unit_labors": unit_labors,
    })


@transaction.atomic
def realization_bom_add(request, project_uuid):
    project = get_object_or_404(Project, uuid=project_uuid)
    if request.method == "POST":
        if project.status == "completed":
            messages.error(request, f"Proyek '{project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin mencatat realisasi material.")
            return redirect("project_realization_detail", uuid=project.uuid)

        bom_item_uuid = request.POST.get("bom_item_uuid", "").strip()
        raw_material_uuid = request.POST.get("raw_material_uuid", "").strip()
        date = request.POST.get("date")
        item_name = request.POST.get("item_name", "").strip()
        unit = request.POST.get("unit", "pcs").strip()
        qty = _clean_decimal(request.POST.get("qty"), Decimal(0))
        unit_cost = _clean_decimal(request.POST.get("unit_cost"), Decimal(0))
        is_substitute = request.POST.get("is_substitute") == "1"
        notes = request.POST.get("notes", "").strip()

        bom_item = None
        raw_material = None

        if bom_item_uuid:
            bom_item = BOMItem.objects.filter(uuid=bom_item_uuid, project=project).first()
            if bom_item:
                if not item_name:
                    item_name = bom_item.name
                if not unit:
                    unit = bom_item.unit
                if bom_item.raw_material and not raw_material_uuid:
                    raw_material = bom_item.raw_material

        if raw_material_uuid:
            raw_material = RawMaterial.objects.filter(uuid=raw_material_uuid).first()
            if raw_material:
                if not item_name:
                    item_name = raw_material.name
                if not unit:
                    unit = raw_material.stock_unit
                if unit_cost == 0:
                    unit_cost = raw_material.last_purchase_price

        if not item_name:
            messages.error(request, "Nama material harus diisi.")
            return redirect("project_realization_detail", uuid=project.uuid)

        if qty <= 0:
            messages.error(request, "Quantity realisasi material harus lebih besar dari 0.")
            return redirect("project_realization_detail", uuid=project.uuid)

        BOMItemRealization.objects.create(
            project=project,
            bom_item=bom_item,
            raw_material=raw_material,
            date=date,
            item_name=item_name,
            unit=unit,
            qty=qty,
            unit_cost=unit_cost,
            is_substitute=is_substitute,
            notes=notes,
        )
        messages.success(request, f"Realisasi material '{item_name}' ({qty} {unit}) berhasil disimpan & mutasi stok dicatat.")
    return redirect("project_realization_detail", uuid=project.uuid)


@transaction.atomic
def realization_bom_delete(request, uuid):
    realization = get_object_or_404(BOMItemRealization, uuid=uuid)
    project_uuid = realization.project.uuid
    if realization.project.status == "completed":
        messages.error(request, f"Proyek '{realization.project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin menghapus log realisasi.")
        return redirect("project_realization_detail", uuid=project_uuid)

    item_name = realization.item_name
    realization.delete()
    messages.success(request, f"Log realisasi material '{item_name}' berhasil dibatalkan dan stok dikembalikan.")
    return redirect("project_realization_detail", uuid=project_uuid)


@transaction.atomic
def realization_labor_add(request, project_uuid):
    project = get_object_or_404(Project, uuid=project_uuid)
    if request.method == "POST":
        if project.status == "completed":
            messages.error(request, f"Proyek '{project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin mencatat realisasi tenaga kerja.")
            return redirect("project_realization_detail", uuid=project.uuid)

        labor_item_uuid = request.POST.get("labor_item_uuid", "").strip()
        date = request.POST.get("date")
        role_name = request.POST.get("role_name", "").strip()
        unit = request.POST.get("unit", "jam").strip()
        quantity = _clean_decimal(request.POST.get("quantity"), Decimal(0))
        rate = _clean_decimal(request.POST.get("rate"), Decimal(0))
        is_additional = request.POST.get("is_additional") == "1"
        notes = request.POST.get("notes", "").strip()

        project_labor = None
        if labor_item_uuid:
            project_labor = ProjectLabor.objects.filter(uuid=labor_item_uuid, project=project).first()
            if project_labor:
                if not role_name:
                    role_name = project_labor.role_name
                if not unit:
                    unit = project_labor.unit

        if not role_name:
            messages.error(request, "Nama/Peran tenaga kerja harus diisi.")
            return redirect("project_realization_detail", uuid=project.uuid)

        if quantity <= 0:
            messages.error(request, "Jumlah/durasi tenaga kerja harus lebih besar dari 0.")
            return redirect("project_realization_detail", uuid=project.uuid)

        LaborRealization.objects.create(
            project=project,
            project_labor=project_labor,
            date=date,
            role_name=role_name,
            unit=unit,
            quantity=quantity,
            rate=rate,
            is_additional=is_additional,
            notes=notes,
        )
        messages.success(request, f"Realisasi tenaga kerja '{role_name}' ({quantity} {unit}) berhasil disimpan.")
    return redirect("project_realization_detail", uuid=project.uuid)


@transaction.atomic
def realization_labor_delete(request, uuid):
    realization = get_object_or_404(LaborRealization, uuid=uuid)
    project_uuid = realization.project.uuid
    if realization.project.status == "completed":
        messages.error(request, f"Proyek '{realization.project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin menghapus log realisasi.")
        return redirect("project_realization_detail", uuid=project_uuid)

    role_name = realization.role_name
    realization.delete()
    messages.success(request, f"Log realisasi tenaga kerja '{role_name}' berhasil dihapus.")
    return redirect("project_realization_detail", uuid=project_uuid)


@transaction.atomic
def realization_fg_add(request, project_uuid):
    project = get_object_or_404(Project, uuid=project_uuid)
    if request.method == "POST":
        if project.status == "completed":
            messages.error(request, f"Proyek '{project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin mencatat realisasi barang jadi.")
            return redirect("project_realization_detail", uuid=project.uuid)

        project_fg_uuid = request.POST.get("project_fg_uuid", "").strip()
        fg_uuid = request.POST.get("fg_uuid", "").strip()
        date = request.POST.get("date")
        quantity = _clean_decimal(request.POST.get("quantity"), Decimal(0))
        unit_cost = _clean_decimal(request.POST.get("unit_cost"), Decimal(0))
        notes = request.POST.get("notes", "").strip()

        project_fg = None
        finished_good = None

        if project_fg_uuid:
            project_fg = ProjectFinishedGood.objects.filter(uuid=project_fg_uuid, project=project).first()
            if project_fg:
                finished_good = project_fg.finished_good
                if unit_cost == 0:
                    unit_cost = project_fg.est_unit_cost or finished_good.standard_cost
        elif fg_uuid:
            finished_good = FinishedGood.objects.filter(uuid=fg_uuid).first()
            if finished_good and unit_cost == 0:
                unit_cost = finished_good.standard_cost

        if not finished_good:
            messages.error(request, "Barang Jadi tidak valid.")
            return redirect("project_realization_detail", uuid=project.uuid)

        if quantity <= 0:
            messages.error(request, "Jumlah pemakaian barang jadi harus lebih besar dari 0.")
            return redirect("project_realization_detail", uuid=project.uuid)

        if finished_good.current_stock < quantity:
            messages.error(request, f"Stock {finished_good.name} tidak cukup (Tersedia: {finished_good.current_stock} {finished_good.unit}).")
            return redirect("project_realization_detail", uuid=project.uuid)

        FinishedGoodRealization.objects.create(
            project=project,
            project_fg=project_fg,
            finished_good=finished_good,
            date=date,
            quantity=quantity,
            unit_cost=unit_cost,
            notes=notes,
        )
        messages.success(request, f"Realisasi {finished_good.name} ({quantity} {finished_good.unit}) berhasil dicatat dan memotong stock.")
    return redirect("project_realization_detail", uuid=project.uuid)


@transaction.atomic
def realization_fg_delete(request, uuid):
    realization = get_object_or_404(FinishedGoodRealization, uuid=uuid)
    project_uuid = realization.project.uuid
    if realization.project.status == "completed":
        messages.error(request, f"Proyek '{realization.project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin menghapus log realisasi.")
        return redirect("project_realization_detail", uuid=project_uuid)

    fg_name = realization.finished_good.name
    realization.delete()
    messages.success(request, f"Log realisasi barang jadi '{fg_name}' dibatalkan dan stock telah dikembalikan.")
    return redirect("project_realization_detail", uuid=project_uuid)


@transaction.atomic
def realization_overhead_add(request, project_uuid):
    project = get_object_or_404(Project, uuid=project_uuid)
    if request.method == "POST":
        if project.status == "completed":
            messages.error(request, f"Proyek '{project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin mencatat realisasi overhead.")
            return redirect("project_realization_detail", uuid=project.uuid)

        overhead_item_uuid = request.POST.get("overhead_item_uuid", "").strip()
        date = request.POST.get("date")
        expense_name = request.POST.get("expense_name", "").strip()
        cost = _clean_decimal(request.POST.get("cost"), Decimal(0))
        is_additional = request.POST.get("is_additional") == "1"
        notes = request.POST.get("notes", "").strip()

        project_overhead = None
        if overhead_item_uuid:
            project_overhead = ProjectOverhead.objects.filter(uuid=overhead_item_uuid, project=project).first()
            if project_overhead and not expense_name:
                expense_name = project_overhead.name

        if not expense_name:
            messages.error(request, "Nama pos pengeluaran overhead harus diisi.")
            return redirect("project_realization_detail", uuid=project.uuid)

        if cost <= 0:
            messages.error(request, "Biaya overhead harus lebih besar dari 0.")
            return redirect("project_realization_detail", uuid=project.uuid)

        OverheadRealization.objects.create(
            project=project,
            project_overhead=project_overhead,
            date=date,
            expense_name=expense_name,
            cost=cost,
            is_additional=is_additional,
            notes=notes,
        )
        messages.success(request, f"Realisasi overhead '{expense_name}' berhasil disimpan.")
    return redirect("project_realization_detail", uuid=project.uuid)


@transaction.atomic
def realization_overhead_delete(request, uuid):
    realization = get_object_or_404(OverheadRealization, uuid=uuid)
    project_uuid = realization.project.uuid
    if realization.project.status == "completed":
        messages.error(request, f"Proyek '{realization.project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin menghapus log realisasi.")
        return redirect("project_realization_detail", uuid=project_uuid)

    expense_name = realization.expense_name
    realization.delete()
    messages.success(request, f"Log realisasi overhead '{expense_name}' berhasil dihapus.")
    return redirect("project_realization_detail", uuid=project_uuid)

