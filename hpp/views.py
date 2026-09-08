from .utils import generate_project_code, get_ordered_bom
from django.shortcuts import render, get_object_or_404, redirect
from django.http import HttpResponse
from django.db import transaction
from django.db.models import Q
from django.contrib import messages
from django.utils import timezone
from decimal import Decimal
from datetime import datetime
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .models import (
    UnitMaster,
    Customer,
    Employee, Attendance, AbsenceType, EmployeeWorkLog,
    Project, BOMItem, ProjectLabor, ProjectOverhead,
    MaterialMaster, LaborMaster, FinishedGood,
    ProjectFinishedGood, FinishedGoodStockMutation,
    RawMaterial, LaborRealization
)


# =========================================================================
# PROJECT VIEWS
# =========================================================================
def project_list(request):
    query = request.GET.get("q", "").strip()
    status = request.GET.get("status", "").strip()
    start_date_from = request.GET.get("start_date_from", "").strip()
    start_date_to = request.GET.get("start_date_to", "").strip()

    projects = Project.objects.all().order_by("-created_at")
    if query:
        projects = projects.filter(Q(name__icontains=query) | Q(code__icontains=query) | Q(customer_name__icontains=query))
    if status:
        projects = projects.filter(status=status)
    if start_date_from:
        projects = projects.filter(start_date__gte=start_date_from)
    if start_date_to:
        projects = projects.filter(start_date__lte=start_date_to)

    return render(request, "hpp/project_list.html", {
        "projects": projects,
        "query": query,
        "selected_status": status,
        "start_date_from": start_date_from,
        "start_date_to": start_date_to,
        "status_choices": Project.STATUS_CHOICES,
    })


@transaction.atomic
def project_create(request):
    if request.method == "POST":
        code = request.POST.get("code", "").strip() or generate_project_code()
        name = request.POST.get("name", "").strip()
        customer_uuid = request.POST.get("customer_uuid", "").strip()
        customer = Customer.objects.filter(uuid=customer_uuid).first() if customer_uuid else None
        manual_customer_name = request.POST.get("customer_name", "").strip()
        customer_name = str(customer) if customer else manual_customer_name

        raw_contract = request.POST.get("contract_value", "0").strip()
        clean_contract = raw_contract.replace("Rp", "").replace("rp", "").replace(" ", "")
        if "," in clean_contract and "." in clean_contract:
            clean_contract = clean_contract.replace(".", "").replace(",", ".")
        elif "." in clean_contract and "," not in clean_contract:
            parts = clean_contract.split(".")
            if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3):
                clean_contract = clean_contract.replace(".", "")
        elif "," in clean_contract:
            clean_contract = clean_contract.replace(",", ".")

        try:
            contract_value = Decimal(clean_contract or 0)
            if contract_value < 0:
                contract_value = Decimal(0)
        except Exception:
            contract_value = Decimal(0)

        try:
            progress_percentage = int(request.POST.get("progress_percentage") or 0)
            progress_percentage = max(0, min(100, progress_percentage))
        except (ValueError, TypeError):
            progress_percentage = 0

        status = request.POST.get("status", "draft")
        start_date = request.POST.get("start_date") or None
        target_date = request.POST.get("target_date") or None
        notes = request.POST.get("notes", "").strip()

        # Validasi: Nama project wajib diisi
        if not name:
            messages.error(request, "Nama project wajib diisi.")
            customers = Customer.objects.all().order_by("name")
            customers_json = json.dumps([
                {
                    "uuid": str(c.uuid),
                    "code": c.code,
                    "name": c.name,
                    "company_name": c.company_name or "",
                    "label": str(c),
                }
                for c in customers
            ])
            return render(request, "hpp/project_form.html", {
                "status_choices": Project.STATUS_CHOICES,
                "auto_code": code or generate_project_code(),
                "customers": customers,
                "customers_json": customers_json,
                "form_data": request.POST,
            })

        # Validasi: Urutan tanggal
        if start_date and target_date:
            try:
                s_d = datetime.strptime(start_date, "%Y-%m-%d").date()
                t_d = datetime.strptime(target_date, "%Y-%m-%d").date()
                if t_d < s_d:
                    messages.error(request, "Target tanggal selesai tidak boleh lebih awal dari tanggal mulai pengerjaan.")
                    customers = Customer.objects.all().order_by("name")
                    customers_json = json.dumps([
                        {
                            "uuid": str(c.uuid),
                            "code": c.code,
                            "name": c.name,
                            "company_name": c.company_name or "",
                            "label": str(c),
                        }
                        for c in customers
                    ])
                    return render(request, "hpp/project_form.html", {
                        "status_choices": Project.STATUS_CHOICES,
                        "auto_code": code or generate_project_code(),
                        "customers": customers,
                        "customers_json": customers_json,
                        "form_data": request.POST,
                    })
            except ValueError:
                pass

        if Project.objects.filter(code=code).exists():
            code = generate_project_code()

        project = Project.objects.create(
            code=code,
            name=name,
            customer=customer,
            customer_name=customer_name,
            contract_value=contract_value,
            progress_percentage=progress_percentage,
            status=status,
            start_date=start_date,
            target_date=target_date,
            notes=notes,
        )
        project.release_stock_if_completed()
        messages.success(request, f"Project '{project.name}' ({project.code}) berhasil dibuat! Silakan lengkapi Bill of Materials (BOM).")
        return redirect("project_detail", uuid=project.uuid)

    auto_code = generate_project_code()
    customers = Customer.objects.all().order_by("name")
    customers_json = json.dumps([
        {
            "uuid": str(c.uuid),
            "code": c.code,
            "name": c.name,
            "company_name": c.company_name or "",
            "label": str(c),
        }
        for c in customers
    ])
    return render(request, "hpp/project_form.html", {
        "status_choices": Project.STATUS_CHOICES,
        "auto_code": auto_code,
        "customers": customers,
        "customers_json": customers_json,
    })


def _clean_decimal(val_str, default=Decimal(0)):
    if val_str is None:
        return default
    if isinstance(val_str, (int, float, Decimal)):
        return Decimal(str(val_str))
    cleaned = str(val_str).replace("Rp", "").replace("rp", "").replace(" ", "").strip()
    if "," in cleaned and "." in cleaned:
        cleaned = cleaned.replace(".", "").replace(",", ".")
    elif "." in cleaned and "," not in cleaned:
        parts = cleaned.split(".")
        if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3):
            cleaned = cleaned.replace(".", "")
    elif "," in cleaned:
        cleaned = cleaned.replace(",", ".")
    try:
        res = Decimal(cleaned or 0)
        return res if res >= 0 else Decimal(0)
    except Exception:
        return default


def project_detail(request, uuid):
    project = get_object_or_404(Project, uuid=uuid)
    materials = get_ordered_bom(project)
    labors = project.labor_items.all().select_related("labor_master")
    overheads = project.overhead_items.all()
    finished_goods = project.finished_good_items.all().select_related("finished_good")
    
    master_materials = MaterialMaster.objects.all().order_by("name")
    raw_materials = RawMaterial.objects.all().order_by("name")
    master_labors = LaborMaster.objects.all().order_by("role_name")
    master_finished_goods = FinishedGood.objects.all().order_by("name")

    unit_materials = UnitMaster.objects.filter(category="raw_material", is_active=True).order_by("name")
    unit_labors = UnitMaster.objects.filter(category="labor", is_active=True).order_by("name")
    unit_fgs = UnitMaster.objects.filter(category="finished_good", is_active=True).order_by("name")
    unit_overheads = UnitMaster.objects.filter(category="overhead", is_active=True).order_by("name")

    raw_materials_json = json.dumps([
        {
            "uuid": str(rm.uuid),
            "code": rm.code,
            "name": rm.name,
            "category": rm.category or "Umum",
            "stock_unit": rm.stock_unit,
            "current_stock": float(rm.current_stock),
            "last_purchase_price": float(rm.last_purchase_price),
            "label": f"[{rm.code}] {rm.name} (Stok: {rm.current_stock:g} {rm.stock_unit} | Rp {rm.last_purchase_price:,.0f})",
        }
        for rm in raw_materials
    ])

    master_labors_json = json.dumps([
        {
            "id": ml.id,
            "uuid": str(ml.uuid),
            "role_name": ml.role_name,
            "unit": ml.unit,
            "standard_rate": float(ml.standard_rate),
            "label": f"{ml.role_name} (Rp {ml.standard_rate:,.0f} / {ml.unit})",
        }
        for ml in master_labors
    ])

    master_finished_goods_json = json.dumps([
        {
            "uuid": str(fg.uuid),
            "sku": fg.sku,
            "name": fg.name,
            "category": fg.category or "Umum",
            "unit": fg.unit,
            "current_stock": float(fg.current_stock),
            "standard_cost": float(fg.standard_cost),
            "label": f"[{fg.sku}] {fg.name} (Stok: {fg.current_stock:g} {fg.unit} | Rp {fg.standard_cost:,.0f})",
        }
        for fg in master_finished_goods
    ])

    return render(request, "hpp/project_detail.html", {
        "project": project,
        "materials": materials,
        "labors": labors,
        "overheads": overheads,
        "finished_goods": finished_goods,
        "master_materials": master_materials,
        "raw_materials": raw_materials,
        "raw_materials_json": raw_materials_json,
        "master_labors": master_labors,
        "master_labors_json": master_labors_json,
        "master_finished_goods": master_finished_goods,
        "master_finished_goods_json": master_finished_goods_json,
        "unit_materials": unit_materials,
        "unit_labors": unit_labors,
        "unit_fgs": unit_fgs,
        "unit_overheads": unit_overheads,
        "status_choices": Project.STATUS_CHOICES,
    })


def project_update(request, uuid):
    project = get_object_or_404(Project, uuid=uuid)
    if request.method == "POST":
        if project.status == "completed":
            messages.error(request, f"Proyek '{project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin mengubah informasi proyek.")
            return redirect("project_detail", uuid=project.uuid)

        new_code = request.POST.get("code", project.code).strip()
        if new_code and new_code != project.code:
            if project.has_realizations:
                messages.warning(
                    request, 
                    f"Kode proyek '{project.code}' tidak dapat diubah karena sudah memiliki riwayat mutasi/realisasi operasional."
                )
            else:
                project.code = new_code

        project.name = request.POST.get("name", project.name).strip()
        customer_uuid = request.POST.get("customer_uuid", "").strip()
        if customer_uuid:
            customer = Customer.objects.filter(uuid=customer_uuid).first()
            if customer:
                project.customer = customer
                project.customer_name = str(customer)
        elif "customer_name" in request.POST:
            project.customer_name = request.POST.get("customer_name", "").strip()

        raw_contract = request.POST.get("contract_value", str(project.contract_value)).strip()
        clean_contract = raw_contract.replace("Rp", "").replace("rp", "").replace(" ", "")
        if "," in clean_contract and "." in clean_contract:
            clean_contract = clean_contract.replace(".", "").replace(",", ".")
        elif "." in clean_contract and "," not in clean_contract:
            parts = clean_contract.split(".")
            if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3):
                clean_contract = clean_contract.replace(".", "")
        elif "," in clean_contract:
            clean_contract = clean_contract.replace(",", ".")
        try:
            project.contract_value = Decimal(clean_contract or 0)
        except Exception:
            pass

        try:
            progress = int(request.POST.get("progress_percentage") or project.progress_percentage)
            project.progress_percentage = max(0, min(100, progress))
        except (ValueError, TypeError):
            pass

        new_status = request.POST.get("status", project.status)
        if new_status == "completed" and project.status != "completed":
            messages.warning(
                request, 
                "Penutupan status ke 'Selesai (Completed)' harus melalui alur resmi tombol 'Closing Project' untuk rekonsiliasi HPP dan Berita Acara (BAP)."
            )
        elif new_status in ["draft", "in_progress", "cancelled"]:
            project.status = new_status

        project.start_date = request.POST.get("start_date") or None
        project.target_date = request.POST.get("target_date") or None
        if "notes" in request.POST:
            project.notes = request.POST.get("notes", "").strip()
        project.save()
        
        messages.success(request, f"Informasi project '{project.name}' berhasil diperbarui.")
    return redirect("project_detail", uuid=project.uuid)


# =========================================================================
# PROJECT CLOSING & REOPEN VIEWS
# =========================================================================
def project_close(request, uuid):
    """
    Eksekusi Penutupan Proyek (Closing Project):
    - Mengubah status ke 'completed' dan progress ke 100%
    - Mencatat completed_date dan closing_notes
    - Mengunci data HPP proyek untuk kepatuhan audit pembukuan
    - Opsional: Melakukan Stock IN ke persediaan gudang jika proyek menghasilkan Barang Jadi
    """
    project = get_object_or_404(Project, uuid=uuid)
    if request.method == "POST":
        if project.status == "completed":
            messages.warning(request, f"Proyek '{project.name}' sudah berstatus Selesai (Closed).")
            return redirect("project_detail", uuid=project.uuid)
        
        date_str = request.POST.get("completed_date")
        if date_str:
            try:
                completed_date = datetime.strptime(date_str, "%Y-%m-%d").date()
            except ValueError:
                completed_date = timezone.now().date()
        else:
            completed_date = timezone.now().date()

        closing_notes = request.POST.get("closing_notes", "").strip()
        store_finished_good = request.POST.get("store_finished_good") in ["1", "on", "true"]

        with transaction.atomic():
            project.status = "completed"
            project.progress_percentage = 100
            project.completed_date = completed_date
            project.closing_notes = closing_notes
            project.save(update_fields=["status", "progress_percentage", "completed_date", "closing_notes", "updated_at"])

            if store_finished_good:
                fg_uuid = request.POST.get("finished_good_uuid", "").strip()
                fg_qty = _clean_decimal(request.POST.get("finished_good_qty", "0"))
                fg_notes = request.POST.get("finished_good_notes", "").strip()

                if fg_uuid and fg_qty > 0:
                    fg = FinishedGood.objects.filter(uuid=fg_uuid).first()
                    if fg:
                        project.release_stock_if_completed(fg=fg, qty=fg_qty, notes=fg_notes)
                        messages.success(
                            request,
                            f"Proyek '{project.name}' ({project.code}) BERHASIL DITUTUP! Stok {fg.name} bertambah +{fg_qty:g} {fg.unit} di gudang."
                        )
                    else:
                        messages.success(request, f"Proyek '{project.name}' ({project.code}) BERHASIL DITUTUP (Status: Selesai, Progress: 100%).")
                else:
                    messages.success(request, f"Proyek '{project.name}' ({project.code}) BERHASIL DITUTUP (Status: Selesai, Progress: 100%).")
            else:
                messages.success(request, f"Proyek '{project.name}' ({project.code}) BERHASIL DITUTUP (Status: Selesai, Progress: 100%). Seluruh data HPP dikunci untuk audit.")

    return redirect("project_detail", uuid=project.uuid)


def project_reopen(request, uuid):
    """
    Membuka kembali proyek yang sudah selesai (Reopen Project):
    - Mengembalikan status ke 'in_progress' dan progress ke 95%
    - Melakukan rollback penambahan stok barang jadi jika sebelumnya di-release
    - Mencatat riwayat alasan pembukaan kembali ke notes
    """
    project = get_object_or_404(Project, uuid=uuid)
    if request.method == "POST":
        if project.status != "completed":
            messages.warning(request, f"Proyek '{project.name}' belum berstatus Selesai, tidak perlu dibuka kembali.")
            return redirect("project_detail", uuid=project.uuid)

        reopen_reason = request.POST.get("reopen_reason", "").strip() or "Revisi data HPP dan realisasi lapangan"
        today_str = timezone.now().strftime("%d/%m/%Y %H:%M")

        with transaction.atomic():
            # Rollback stok barang jadi jika pernah dirilis
            if project.stock_released:
                project.revert_stock_release()

            project.status = "in_progress"
            project.progress_percentage = 95
            audit_entry = f"\n[BUKA KEMBALI {today_str}] Alasan: {reopen_reason}"
            project.notes = (project.notes or "") + audit_entry
            project.save(update_fields=["status", "progress_percentage", "notes", "updated_at"])

            messages.success(request, f"Proyek '{project.name}' ({project.code}) berhasil dibuka kembali. Anda dapat menambahkan atau mengoreksi data.")

    return redirect("project_detail", uuid=project.uuid)


def project_closing_bap(request, uuid):
    """
    Mencetak Berita Acara Penyelesaian Proyek (BAP) resmi.
    """
    project = get_object_or_404(Project, uuid=uuid)
    materials = get_ordered_bom(project)
    labors = project.labor_items.all().select_related("labor_master")
    overheads = project.overhead_items.all()
    finished_goods = project.finished_good_items.all().select_related("finished_good")

    # Ambil mutasi stok hasil produksi jika ada
    prod_mutations = project.stock_mutations.filter(mutation_type="IN", reference_no=f"PROD-{project.code}").select_related("finished_good")

    return render(request, "hpp/project_closing_bap.html", {
        "project": project,
        "materials": materials,
        "labors": labors,
        "overheads": overheads,
        "finished_goods": finished_goods,
        "prod_mutations": prod_mutations,
        "now": timezone.now(),
    })


# =========================================================================
# BOM CRUD
# =========================================================================
def bom_item_add(request, project_uuid):
    project = get_object_or_404(Project, uuid=project_uuid)
    if request.method == "POST":
        if project.status == "completed":
            messages.error(request, f"Proyek '{project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin menambah item BOM.")
            return redirect("project_detail", uuid=project.uuid)

        name = request.POST.get("name", "").strip()
        item_type = request.POST.get("item_type", "material")
        raw_material_uuid = request.POST.get("raw_material_uuid", "").strip()
        raw_material = RawMaterial.objects.filter(uuid=raw_material_uuid).first() if raw_material_uuid else None
        if raw_material and not name:
            name = raw_material.name

        unit = request.POST.get("unit", "").strip()
        if not unit and raw_material:
            unit = raw_material.stock_unit
        if not unit:
            unit = "pcs"

        raw_est_qty = request.POST.get("est_qty", "1").strip().replace(",", ".")
        try:
            est_qty = Decimal(raw_est_qty or 1)
            if est_qty < 0:
                est_qty = Decimal(0)
        except Exception:
            est_qty = Decimal(1)

        raw_est_cost = request.POST.get("est_unit_cost", "0").strip()
        clean_cost = raw_est_cost.replace("Rp", "").replace("rp", "").replace(" ", "")
        if "," in clean_cost and "." in clean_cost:
            clean_cost = clean_cost.replace(".", "").replace(",", ".")
        elif "." in clean_cost and "," not in clean_cost:
            parts = clean_cost.split(".")
            if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3):
                clean_cost = clean_cost.replace(".", "")
        elif "," in clean_cost:
            clean_cost = clean_cost.replace(",", ".")
        try:
            est_unit_cost = Decimal(clean_cost or 0)
            if est_unit_cost < 0:
                est_unit_cost = Decimal(0)
        except Exception:
            est_unit_cost = Decimal(0)

        if est_unit_cost == 0 and raw_material:
            est_unit_cost = raw_material.last_purchase_price

        parent_uuid = request.POST.get("parent_uuid") or None
        parent = BOMItem.objects.filter(uuid=parent_uuid, project=project).first() if parent_uuid else None
        notes = request.POST.get("notes", "").strip()

        if not name:
            messages.error(request, "Nama item BOM wajib diisi.")
            return redirect("project_detail", uuid=project.uuid)

        BOMItem.objects.create(
            project=project,
            parent=parent,
            raw_material=raw_material,
            name=name,
            item_type=item_type,
            unit=unit,
            est_qty=est_qty,
            est_unit_cost=est_unit_cost,
            notes=notes,
        )
        messages.success(request, f"Item BOM '{name}' berhasil ditambahkan.")
    return redirect("project_detail", uuid=project.uuid)


def bom_item_update(request, uuid):
    item = get_object_or_404(BOMItem, uuid=uuid)
    if request.method == "POST":
        if item.project.status == "completed":
            messages.error(request, f"Proyek '{item.project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin mengedit item BOM.")
            return redirect("project_detail", uuid=item.project.uuid)

        name = request.POST.get("name", item.name).strip()
        item_type = request.POST.get("item_type", item.item_type)
        raw_material_uuid = request.POST.get("raw_material_uuid", "").strip()
        raw_material = RawMaterial.objects.filter(uuid=raw_material_uuid).first() if raw_material_uuid else item.raw_material

        unit = request.POST.get("unit", item.unit).strip()

        raw_est_qty = request.POST.get("est_qty", str(item.est_qty)).strip().replace(",", ".")
        try:
            est_qty = Decimal(raw_est_qty or 0)
            if est_qty < 0:
                est_qty = Decimal(0)
        except Exception:
            est_qty = item.est_qty

        raw_est_cost = request.POST.get("est_unit_cost", str(item.est_unit_cost)).strip()
        clean_cost = raw_est_cost.replace("Rp", "").replace("rp", "").replace(" ", "")
        if "," in clean_cost and "." in clean_cost:
            clean_cost = clean_cost.replace(".", "").replace(",", ".")
        elif "." in clean_cost and "," not in clean_cost:
            parts = clean_cost.split(".")
            if len(parts) > 2 or (len(parts) == 2 and len(parts[1]) == 3):
                clean_cost = clean_cost.replace(".", "")
        elif "," in clean_cost:
            clean_cost = clean_cost.replace(",", ".")
        try:
            est_unit_cost = Decimal(clean_cost or 0)
            if est_unit_cost < 0:
                est_unit_cost = Decimal(0)
        except Exception:
            est_unit_cost = item.est_unit_cost

        parent_uuid = request.POST.get("parent_uuid") or None
        parent = BOMItem.objects.filter(uuid=parent_uuid, project=item.project).exclude(id=item.id).first() if parent_uuid else None
        notes = request.POST.get("notes", item.notes).strip()

        if name:
            item.name = name
        item.item_type = item_type
        item.raw_material = raw_material
        item.unit = unit
        item.est_qty = est_qty
        item.est_unit_cost = est_unit_cost
        item.parent = parent
        item.notes = notes
        item.save()
        messages.success(request, f"Item BOM '{item.name}' berhasil diperbarui.")
    return redirect("project_detail", uuid=item.project.uuid)


def bom_item_delete(request, uuid):
    item = get_object_or_404(BOMItem, uuid=uuid)
    project_uuid = item.project.uuid
    if request.method == "POST":
        if item.project.status == "completed":
            messages.error(request, f"Proyek '{item.project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin menghapus item BOM.")
            return redirect("project_detail", uuid=project_uuid)

        if item.children.exists():
            messages.error(request, f"Item '{item.name}' tidak dapat dihapus karena masih memiliki sub-item / komponen di dalamnya. Hapus atau pindahkan sub-item terlebih dahulu.")
            return redirect("project_detail", uuid=project_uuid)

        if item.realizations.exists():
            messages.error(request, f"Item '{item.name}' tidak dapat dihapus karena sudah memiliki riwayat transaksi realisasi pengeluaran.")
            return redirect("project_detail", uuid=project_uuid)

        name = item.name
        item.delete()
        messages.success(request, f"Item BOM '{name}' berhasil dihapus.")
    return redirect("project_detail", uuid=project_uuid)


# =========================================================================
# LABOR CRUD
# =========================================================================
@transaction.atomic
def labor_add(request, project_uuid):
    project = get_object_or_404(Project, uuid=project_uuid)
    if request.method == "POST":
        if project.status == "completed":
            messages.error(request, f"Proyek '{project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin menambah data tenaga kerja.")
            return redirect("project_detail", uuid=project.uuid)

        role_name = request.POST.get("role_name", "").strip()
        unit = request.POST.get("unit", "jam").strip()
        est_quantity = _clean_decimal(request.POST.get("est_quantity"), Decimal(1))
        est_rate = _clean_decimal(request.POST.get("est_rate"), Decimal(0))
        act_quantity = _clean_decimal(request.POST.get("act_quantity"), Decimal(0))
        act_rate = _clean_decimal(request.POST.get("act_rate"), Decimal(0))
        notes = request.POST.get("notes", "").strip()

        if not role_name:
            messages.error(request, "Peran / posisi tenaga kerja wajib diisi.")
            return redirect("project_detail", uuid=project.uuid)

        labor = ProjectLabor.objects.create(
            project=project,
            role_name=role_name,
            unit=unit,
            est_quantity=est_quantity,
            est_rate=est_rate,
            act_quantity=act_quantity,
            act_rate=act_rate,
            notes=notes,
        )
        messages.success(request, f"Tenaga kerja '{labor.role_name}' ({est_quantity:g} {unit}) berhasil ditambahkan ke estimasi HPP.")
    return redirect("project_detail", uuid=project.uuid)


@transaction.atomic
def labor_update(request, uuid):
    item = get_object_or_404(ProjectLabor, uuid=uuid)
    if request.method == "POST":
        if item.project.status == "completed":
            messages.error(request, f"Proyek '{item.project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin mengedit data tenaga kerja.")
            return redirect("project_detail", uuid=item.project.uuid)

        role_name = request.POST.get("role_name", "").strip()
        if role_name:
            item.role_name = role_name
        item.unit = request.POST.get("unit", item.unit).strip()
        item.est_quantity = _clean_decimal(request.POST.get("est_quantity"), item.est_quantity)
        item.est_rate = _clean_decimal(request.POST.get("est_rate"), item.est_rate)
        item.act_quantity = _clean_decimal(request.POST.get("act_quantity"), item.act_quantity)
        item.act_rate = _clean_decimal(request.POST.get("act_rate"), item.act_rate)
        item.notes = request.POST.get("notes", item.notes).strip()
        item.save()
        messages.success(request, f"Data tenaga kerja '{item.role_name}' berhasil diperbarui.")
    return redirect("project_detail", uuid=item.project.uuid)


def labor_delete(request, uuid):
    item = get_object_or_404(ProjectLabor, uuid=uuid)
    project_uuid = item.project.uuid
    if request.method == "POST":
        if item.project.status == "completed":
            messages.error(request, f"Proyek '{item.project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin menghapus data tenaga kerja.")
            return redirect("project_detail", uuid=project_uuid)

        if item.realizations.exists():
            messages.error(request, f"Tenaga kerja '{item.role_name}' tidak dapat dihapus karena sudah memiliki riwayat realisasi.")
            return redirect("project_detail", uuid=project_uuid)
        role = item.role_name
        item.delete()
        messages.success(request, f"Tenaga kerja '{role}' berhasil dihapus.")
    return redirect("project_detail", uuid=project_uuid)


# =========================================================================
# OVERHEAD CRUD
# =========================================================================
@transaction.atomic
def overhead_add(request, project_uuid):
    project = get_object_or_404(Project, uuid=project_uuid)
    if request.method == "POST":
        if project.status == "completed":
            messages.error(request, f"Proyek '{project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin menambah biaya overhead.")
            return redirect("project_detail", uuid=project.uuid)

        name = request.POST.get("name", "").strip()
        est_cost = _clean_decimal(request.POST.get("est_cost"), Decimal(0))
        act_cost = _clean_decimal(request.POST.get("act_cost"), Decimal(0))
        notes = request.POST.get("notes", "").strip()

        if not name:
            messages.error(request, "Deskripsi pos biaya overhead wajib diisi.")
            return redirect("project_detail", uuid=project.uuid)

        ProjectOverhead.objects.create(
            project=project,
            name=name,
            est_cost=est_cost,
            act_cost=act_cost,
            notes=notes,
        )
        messages.success(request, f"Biaya overhead '{name}' (Rp {est_cost:,.0f}) berhasil ditambahkan ke estimasi HPP.")
    return redirect("project_detail", uuid=project.uuid)


@transaction.atomic
def overhead_update(request, uuid):
    item = get_object_or_404(ProjectOverhead, uuid=uuid)
    if request.method == "POST":
        if item.project.status == "completed":
            messages.error(request, f"Proyek '{item.project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin mengedit biaya overhead.")
            return redirect("project_detail", uuid=item.project.uuid)

        name = request.POST.get("name", "").strip()
        if name:
            item.name = name
        item.est_cost = _clean_decimal(request.POST.get("est_cost"), item.est_cost)
        item.act_cost = _clean_decimal(request.POST.get("act_cost"), item.act_cost)
        item.notes = request.POST.get("notes", item.notes).strip()
        item.save()
        messages.success(request, f"Biaya overhead '{item.name}' berhasil diperbarui.")
    return redirect("project_detail", uuid=item.project.uuid)


def overhead_delete(request, uuid):
    item = get_object_or_404(ProjectOverhead, uuid=uuid)
    project_uuid = item.project.uuid
    if request.method == "POST":
        if item.project.status == "completed":
            messages.error(request, f"Proyek '{item.project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin menghapus biaya overhead.")
            return redirect("project_detail", uuid=project_uuid)

        if item.realizations.exists():
            messages.error(request, f"Biaya overhead '{item.name}' tidak dapat dihapus karena sudah memiliki riwayat realisasi.")
            return redirect("project_detail", uuid=project_uuid)
        name = item.name
        item.delete()
        messages.success(request, f"Biaya overhead '{name}' berhasil dihapus.")
    return redirect("project_detail", uuid=project_uuid)


# =========================================================================
# PROJECT FINISHED GOODS (Barang Jadi Penyusun Project)
# =========================================================================
@transaction.atomic
def project_fg_add(request, project_uuid):
    project = get_object_or_404(Project, uuid=project_uuid)
    if request.method == "POST":
        if project.status == "completed":
            messages.error(request, f"Proyek '{project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin menambah barang jadi.")
            return redirect("project_detail", uuid=project.uuid)

        fg_uuid = request.POST.get("finished_good_uuid", "").strip()
        fg = FinishedGood.objects.filter(uuid=fg_uuid).first() if fg_uuid else None
        if not fg:
            messages.error(request, "Barang jadi yang dipilih tidak valid.")
            return redirect("project_detail", uuid=project.uuid)

        est_qty = _clean_decimal(request.POST.get("est_qty"), Decimal(1))
        est_unit_cost = _clean_decimal(request.POST.get("est_unit_cost"), fg.standard_cost)
        act_qty = _clean_decimal(request.POST.get("act_qty"), Decimal(0))
        act_unit_cost = _clean_decimal(request.POST.get("act_unit_cost"), fg.standard_cost)
        notes = request.POST.get("notes", "").strip()

        if est_qty <= 0:
            messages.error(request, "Kuantitas estimasi barang jadi harus lebih besar dari 0.")
            return redirect("project_detail", uuid=project.uuid)

        # Cek apakah sudah ada barang jadi yang sama di project ini
        existing = ProjectFinishedGood.objects.filter(project=project, finished_good=fg).first()
        if existing:
            existing.est_qty += est_qty
            if est_unit_cost > 0:
                existing.est_unit_cost = est_unit_cost
            if notes:
                existing.notes = f"{existing.notes}; {notes}".strip("; ")
            existing.save()
            messages.success(request, f"Kuantitas barang jadi '{fg.name}' diperbarui menjadi {existing.est_qty:g} {fg.unit}.")
        else:
            ProjectFinishedGood.objects.create(
                project=project,
                finished_good=fg,
                est_qty=est_qty,
                est_unit_cost=est_unit_cost,
                act_qty=act_qty,
                act_unit_cost=act_unit_cost,
                notes=notes,
            )
            messages.success(request, f"Barang jadi '{fg.name}' ({est_qty:g} {fg.unit}) berhasil ditambahkan ke project.")
    return redirect("project_detail", uuid=project.uuid)


@transaction.atomic
def project_fg_update(request, uuid):
    item = get_object_or_404(ProjectFinishedGood, uuid=uuid)
    if request.method == "POST":
        if item.project.status == "completed":
            messages.error(request, f"Proyek '{item.project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin mengedit barang jadi.")
            return redirect("project_detail", uuid=item.project.uuid)

        fg_uuid = request.POST.get("finished_good_uuid", "").strip()
        if fg_uuid:
            fg = FinishedGood.objects.filter(uuid=fg_uuid).first()
            if fg:
                item.finished_good = fg
        item.est_qty = _clean_decimal(request.POST.get("est_qty"), item.est_qty)
        item.est_unit_cost = _clean_decimal(request.POST.get("est_unit_cost"), item.est_unit_cost)
        item.act_qty = _clean_decimal(request.POST.get("act_qty"), item.act_qty)
        item.act_unit_cost = _clean_decimal(request.POST.get("act_unit_cost"), item.act_unit_cost)
        item.notes = request.POST.get("notes", item.notes).strip()
        item.save()
        messages.success(request, f"Data barang jadi '{item.finished_good.name}' berhasil diperbarui.")
    return redirect("project_detail", uuid=item.project.uuid)


def project_fg_delete(request, uuid):
    item = get_object_or_404(ProjectFinishedGood, uuid=uuid)
    project_uuid = item.project.uuid
    if request.method == "POST":
        if item.project.status == "completed":
            messages.error(request, f"Proyek '{item.project.code}' sudah ditutup dan terkunci. Harap buka kembali proyek jika ingin menghapus barang jadi.")
            return redirect("project_detail", uuid=project_uuid)

        if item.realizations.exists():
            messages.error(request, f"Barang jadi '{item.finished_good.name}' tidak dapat dihapus karena sudah ada riwayat pemakaian riil.")
            return redirect("project_detail", uuid=project_uuid)
        name = item.finished_good.name
        item.delete()
        messages.success(request, f"Barang jadi '{name}' berhasil dihapus dari project.")
    return redirect("project_detail", uuid=project_uuid)


# =========================================================================
# STOCK MANAGEMENT & MUTASI BARANG JADI
# =========================================================================
def stock_report(request):
    """
    Laporan Rekap Stock Semua Barang Jadi
    """
    query = request.GET.get("q", "").strip()
    finished_goods = FinishedGood.objects.all().order_by("sku")
    if query:
        finished_goods = finished_goods.filter(Q(name__icontains=query) | Q(sku__icontains=query))

    total_units = sum(fg.current_stock for fg in finished_goods)
    total_valuation = sum(fg.current_stock * fg.standard_cost for fg in finished_goods)

    for fg in finished_goods:
        fg.stock_valuation = fg.current_stock * fg.standard_cost
        fg.has_transactions = fg.mutations.exists() or fg.project_usages.exists()

    return render(request, "hpp/stock_report.html", {
        "finished_goods": finished_goods,
        "total_units": total_units,
        "total_valuation": total_valuation,
        "query": query,
    })


def finished_good_create(request):
    if request.method == "POST":
        sku = request.POST.get("sku", "").strip()
        name = request.POST.get("name", "").strip()
        unit = request.POST.get("unit", "unit").strip()
        standard_cost = Decimal(request.POST.get("standard_cost") or 0)
        initial_stock = Decimal(request.POST.get("initial_stock") or 0)
        notes = request.POST.get("notes", "")

        fg = FinishedGood.objects.create(
            sku=sku,
            name=name,
            unit=unit,
            standard_cost=standard_cost,
            current_stock=initial_stock,
            notes=notes,
        )

        if initial_stock > 0:
            FinishedGoodStockMutation.objects.create(
                finished_good=fg,
                mutation_type="IN",
                quantity=initial_stock,
                balance_after=initial_stock,
                reference_no="INITIAL-STOCK",
                notes="Saldo awal barang jadi"
            )

        return redirect("stock_report")
    return render(request, "hpp/finished_good_form.html")


def stock_card(request, uuid):
    """
    Kartu Stock per Barang Jadi -> Delegasi ke views_finished_good
    """
    return views_finished_good.finished_good_stock_card(request, uuid)


def stock_mutation_create(request, uuid):
    """
    Manual Stock IN / OUT (Penyesuaian / Restock Pabrik)
    """
    fg = get_object_or_404(FinishedGood, uuid=uuid)
    if request.method == "POST":
        mutation_type = request.POST.get("mutation_type")  # IN / OUT
        quantity = Decimal(request.POST.get("quantity") or 0)
        reference_no = request.POST.get("reference_no", "").strip()
        notes = request.POST.get("notes", "").strip()

        if quantity > 0:
            if mutation_type == "IN":
                fg.current_stock += quantity
            elif mutation_type == "OUT":
                fg.current_stock -= quantity
            fg.save()

            FinishedGoodStockMutation.objects.create(
                finished_good=fg,
                mutation_type=mutation_type,
                quantity=quantity,
                balance_after=fg.current_stock,
                reference_no=reference_no or "MANUAL-ADJ",
                notes=notes,
            )

    return redirect("stock_card", uuid=fg.uuid)


# =========================================================================
# EXPORT & PRINT
# =========================================================================
def export_project_excel(request, uuid):
    project = get_object_or_404(Project, uuid=uuid)
    wb = openpyxl.Workbook()

    # Styling definitions
    font_family = "Calibri"
    thin_side = Side(border_style="thin", color="CBD5E1")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    double_bottom_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=Side(border_style="double", color="0F172A"))
    
    title_font = Font(name=font_family, size=14, bold=True, color="1E1B4B")
    section_title_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    section_title_fill = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
    
    tbl_header_font = Font(name=font_family, size=10, bold=True, color="1E293B")
    tbl_header_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    
    sub_header_font = Font(name=font_family, size=10, bold=True, color="1E1B4B")
    sub_header_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")

    total_font = Font(name=font_family, size=10, bold=True, color="0F172A")
    total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    badge_saving_font = Font(name=font_family, size=9, bold=True, color="047857")
    badge_saving_fill = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    badge_over_font = Font(name=font_family, size=9, bold=True, color="B91C1C")
    badge_over_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    badge_neutral_font = Font(name=font_family, size=9, bold=True, color="475569")
    badge_neutral_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    currency_fmt = '"Rp" #,##0;("Rp" #,##0);"-"'
    percent_fmt = '0.00%'
    qty_fmt = '#,##0.00'

    def format_row(ws, row_idx, font=None, fill=None, border=None, alignment=None):
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if font: cell.font = font
            if fill: cell.fill = fill
            if border: cell.border = border
            if alignment: cell.alignment = alignment

    def autofit(ws):
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                if cell.value is not None:
                    val_str = str(cell.value)
                    if "\n" in val_str:
                        val_str = max(val_str.split("\n"), key=len)
                    if len(val_str) > max_len:
                        max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)

    # -------------------------------------------------------------------------
    # SHEET 1: RINGKASAN & KPI
    # -------------------------------------------------------------------------
    ws_kpi = wb.active
    ws_kpi.title = "Ringkasan & KPI"
    ws_kpi.views.sheetView[0].showGridLines = True

    ws_kpi.append(["LAPORAN EKSEKUTIF HPP & KINERJA PROJECT"])
    ws_kpi["A1"].font = title_font
    ws_kpi.append([])

    # Metadata Project
    ws_kpi.append(["KODE PROJECT", project.code, "", "STATUS PROJECT", project.get_status_display().upper()])
    ws_kpi.append(["NAMA PROJECT", project.name, "", "PROGRESS FISIK", f"{project.progress_percentage}%"])
    ws_kpi.append(["CUSTOMER / KLIEN", project.customer_name or "-", "", "TANGGAL MULAI", str(project.start_date or "-")])
    ws_kpi.append(["NILAI KONTRAK", float(project.contract_value), "", "TARGET SELESAI", str(project.target_date or "-")])
    ws_kpi.cell(row=6, column=2).number_format = currency_fmt
    
    for r in range(3, 7):
        ws_kpi.cell(row=r, column=1).font = Font(name=font_family, size=10, bold=True, color="475569")
        ws_kpi.cell(row=r, column=2).font = Font(name=font_family, size=10, bold=True)
        ws_kpi.cell(row=r, column=4).font = Font(name=font_family, size=10, bold=True, color="475569")
        ws_kpi.cell(row=r, column=5).font = Font(name=font_family, size=10, bold=True)
    
    ws_kpi.append([])

    # KPI Table
    ws_kpi.append(["KOMPONEN BIAYA & LABA", "ESTIMASI (RENCANA)", "REALISASI (AKTUAL)", "DEVIASI (SELISIH)", "RASIO DEVIASI", "STATUS EVALUASI"])
    kpi_hdr_row = ws_kpi.max_row
    for c in range(1, 7):
        cell = ws_kpi.cell(row=kpi_hdr_row, column=c)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c >= 2 else "left", vertical="center")

    kpi_data = [
        ("1. Biaya Material (BOM)", float(project.total_est_material), float(project.total_act_material), True),
        ("2. Biaya Tenaga Kerja (Labor)", float(project.total_est_labor), float(project.total_act_labor), True),
        ("3. Biaya Barang Jadi (FG)", float(project.total_est_finished_goods), float(project.total_act_finished_goods), True),
        ("4. Biaya Overhead & Lainnya", float(project.total_est_overhead), float(project.total_act_overhead), True),
        ("TOTAL BIAYA HPP", float(project.total_hpp_estimated), float(project.total_hpp_actual), True),
        ("Nilai Kontrak / Penjualan", float(project.contract_value), float(project.contract_value), False),
        ("Gross Profit (Laba Kotor)", float(project.est_gross_profit), float(project.act_gross_profit), False),
    ]

    for label, est, act, is_cost in kpi_data:
        diff = act - est
        ratio = (diff / est) if est > 0 else 0.0
        
        if is_cost:
            if diff < 0:
                status_text, b_font, b_fill = "HEMAT (SAVING)", badge_saving_font, badge_saving_fill
            elif diff > 0:
                status_text, b_font, b_fill = "LEBIH (OVER BUDGET)", badge_over_font, badge_over_fill
            else:
                status_text, b_font, b_fill = "SESUAI BUDGET", badge_neutral_font, badge_neutral_fill
        else:
            if "Gross Profit" in label:
                if diff > 0:
                    status_text, b_font, b_fill = "PROFIT NAIK", badge_saving_font, badge_saving_fill
                elif diff < 0:
                    status_text, b_font, b_fill = "PROFIT TURUN", badge_over_font, badge_over_fill
                else:
                    status_text, b_font, b_fill = "SESUAI TARGET", badge_neutral_font, badge_neutral_fill
            else:
                status_text, b_font, b_fill = "-", badge_neutral_font, badge_neutral_fill

        ws_kpi.append([label, est, act, diff, ratio, status_text])
        r_idx = ws_kpi.max_row
        
        is_highlight = label in ["TOTAL BIAYA HPP", "Gross Profit (Laba Kotor)"]
        for c in range(1, 7):
            cell = ws_kpi.cell(row=r_idx, column=c)
            cell.border = double_bottom_border if is_highlight else thin_border
            if is_highlight:
                cell.font = total_font
                cell.fill = total_fill
            if c in [2, 3, 4]:
                cell.number_format = currency_fmt
                cell.alignment = Alignment(horizontal="right")
            elif c == 5:
                cell.number_format = percent_fmt
                cell.alignment = Alignment(horizontal="right")
            elif c == 6:
                cell.alignment = Alignment(horizontal="center")
                cell.font = b_font
                cell.fill = b_fill

    # Margin Row
    est_margin = float(project.est_margin_pct) / 100.0
    act_margin = float(project.act_margin_pct) / 100.0
    margin_diff = act_margin - est_margin
    ws_kpi.append(["Margin Keuntungan (%)", est_margin, act_margin, margin_diff, "", ""])
    r_margin = ws_kpi.max_row
    for c in range(1, 7):
        cell = ws_kpi.cell(row=r_margin, column=c)
        cell.font = total_font
        cell.fill = total_fill
        cell.border = double_bottom_border
        if c in [2, 3, 4]:
            cell.number_format = percent_fmt
            cell.alignment = Alignment(horizontal="right")

    autofit(ws_kpi)

    # -------------------------------------------------------------------------
    # SHEET 2: RENCANA HPP (ESTIMASI)
    # -------------------------------------------------------------------------
    ws_plan = wb.create_sheet(title="Rencana HPP (Estimasi)")
    ws_plan.views.sheetView[0].showGridLines = True

    ws_plan.append([f"ANGGARAN & RENCANA BIAYA HPP - {project.code}"])
    ws_plan["A1"].font = title_font
    ws_plan.append([])

    # 1. BOM Materials
    ws_plan.append(["1. BILL OF MATERIALS (BOM) & BAHAN BAKU"])
    ws_plan.cell(row=ws_plan.max_row, column=1).font = section_title_font
    ws_plan.cell(row=ws_plan.max_row, column=1).fill = section_title_fill
    ws_plan.append(["No", "Tipe", "Item / Material", "Satuan", "Est. Kuantitas", "Est. Harga Satuan", "Est. Total Biaya", "Catatan"])
    hdr_row = ws_plan.max_row
    for c in range(1, 9):
        cell = ws_plan.cell(row=hdr_row, column=c)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c in [1, 2, 4] else ("right" if c in [5, 6, 7] else "left"))

    no = 1
    for m in project.bom_items.all().order_by("parent__id", "id"):
        indent = "   " if m.parent else ""
        ws_plan.append([
            no,
            m.get_item_type_display(),
            f"{indent}{m.name}",
            m.unit,
            float(m.est_qty),
            float(m.est_unit_cost),
            float(m.est_total),
            m.notes or "-"
        ])
        curr_row = ws_plan.max_row
        for c in range(1, 9):
            cell = ws_plan.cell(row=curr_row, column=c)
            cell.border = thin_border
            if c in [5]: cell.number_format = qty_fmt; cell.alignment = Alignment(horizontal="right")
            elif c in [6, 7]: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")
            elif c in [1, 2, 4]: cell.alignment = Alignment(horizontal="center")
        no += 1

    ws_plan.append(["", "SUBTOTAL MATERIAL", "", "", "", "", float(project.total_est_material), ""])
    sub_row = ws_plan.max_row
    for c in range(1, 9):
        cell = ws_plan.cell(row=sub_row, column=c)
        cell.font = total_font
        cell.fill = sub_header_fill
        cell.border = double_bottom_border
        if c == 7: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")

    ws_plan.append([])

    # 2. Labor
    ws_plan.append(["2. ESTIMASI BIAYA TENAGA KERJA (LABOR)"])
    ws_plan.cell(row=ws_plan.max_row, column=1).font = section_title_font
    ws_plan.cell(row=ws_plan.max_row, column=1).fill = section_title_fill
    ws_plan.append(["No", "Peran / Posisi Tenaga Kerja", "Satuan", "Est. Volume", "Est. Tarif Satuan", "Est. Total Biaya"])
    hdr_row = ws_plan.max_row
    for c in range(1, 7):
        cell = ws_plan.cell(row=hdr_row, column=c)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c in [1, 3] else ("right" if c in [4, 5, 6] else "left"))

    no = 1
    for l in project.labor_items.all():
        ws_plan.append([no, l.role_name, l.unit, float(l.est_quantity), float(l.est_rate), float(l.est_total)])
        curr_row = ws_plan.max_row
        for c in range(1, 7):
            cell = ws_plan.cell(row=curr_row, column=c)
            cell.border = thin_border
            if c == 4: cell.number_format = qty_fmt; cell.alignment = Alignment(horizontal="right")
            elif c in [5, 6]: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")
            elif c in [1, 3]: cell.alignment = Alignment(horizontal="center")
        no += 1

    ws_plan.append(["", "SUBTOTAL TENAGA KERJA", "", "", "", float(project.total_est_labor)])
    sub_row = ws_plan.max_row
    for c in range(1, 7):
        cell = ws_plan.cell(row=sub_row, column=c)
        cell.font = total_font
        cell.fill = sub_header_fill
        cell.border = double_bottom_border
        if c == 6: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")

    ws_plan.append([])

    # 3. Finished Goods
    ws_plan.append(["3. ESTIMASI BARANG JADI (FINISHED GOODS)"])
    ws_plan.cell(row=ws_plan.max_row, column=1).font = section_title_font
    ws_plan.cell(row=ws_plan.max_row, column=1).fill = section_title_fill
    ws_plan.append(["No", "SKU", "Nama Barang Jadi", "Satuan", "Est. Kuantitas", "Est. Harga Satuan", "Est. Total Biaya", "Catatan"])
    hdr_row = ws_plan.max_row
    for c in range(1, 9):
        cell = ws_plan.cell(row=hdr_row, column=c)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c in [1, 2, 4] else ("right" if c in [5, 6, 7] else "left"))

    no = 1
    for fg in project.finished_good_items.all():
        ws_plan.append([
            no,
            fg.finished_good.sku,
            fg.finished_good.name,
            fg.finished_good.unit,
            float(fg.est_qty),
            float(fg.est_unit_cost),
            float(fg.est_total),
            fg.notes or "-"
        ])
        curr_row = ws_plan.max_row
        for c in range(1, 9):
            cell = ws_plan.cell(row=curr_row, column=c)
            cell.border = thin_border
            if c == 5: cell.number_format = qty_fmt; cell.alignment = Alignment(horizontal="right")
            elif c in [6, 7]: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")
            elif c in [1, 2, 4]: cell.alignment = Alignment(horizontal="center")
        no += 1

    ws_plan.append(["", "SUBTOTAL BARANG JADI", "", "", "", "", float(project.total_est_finished_goods), ""])
    sub_row = ws_plan.max_row
    for c in range(1, 9):
        cell = ws_plan.cell(row=sub_row, column=c)
        cell.font = total_font
        cell.fill = sub_header_fill
        cell.border = double_bottom_border
        if c == 7: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")

    ws_plan.append([])

    # 4. Overhead
    ws_plan.append(["4. ESTIMASI BIAYA OVERHEAD & OPERASIONAL"])
    ws_plan.cell(row=ws_plan.max_row, column=1).font = section_title_font
    ws_plan.cell(row=ws_plan.max_row, column=1).fill = section_title_fill
    ws_plan.append(["No", "Deskripsi Overhead / Operasional", "Est. Biaya"])
    hdr_row = ws_plan.max_row
    for c in range(1, 4):
        cell = ws_plan.cell(row=hdr_row, column=c)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c == 1 else ("right" if c == 3 else "left"))

    no = 1
    for o in project.overhead_items.all():
        ws_plan.append([no, o.name, float(o.est_cost)])
        curr_row = ws_plan.max_row
        for c in range(1, 4):
            cell = ws_plan.cell(row=curr_row, column=c)
            cell.border = thin_border
            if c == 3: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")
            elif c == 1: cell.alignment = Alignment(horizontal="center")
        no += 1

    ws_plan.append(["", "SUBTOTAL OVERHEAD", float(project.total_est_overhead)])
    sub_row = ws_plan.max_row
    for c in range(1, 4):
        cell = ws_plan.cell(row=sub_row, column=c)
        cell.font = total_font
        cell.fill = sub_header_fill
        cell.border = double_bottom_border
        if c == 3: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")

    ws_plan.append([])
    ws_plan.append(["GRAND TOTAL ESTIMASI HPP PROJECT", "", "", "", "", "", float(project.total_hpp_estimated), ""])
    grand_row = ws_plan.max_row
    for c in range(1, 9):
        cell = ws_plan.cell(row=grand_row, column=c)
        cell.font = Font(name=font_family, size=11, bold=True, color="1E1B4B")
        cell.fill = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
        cell.border = double_bottom_border
        if c == 7: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")

    autofit(ws_plan)

    # -------------------------------------------------------------------------
    # SHEET 3: REALISASI LAPANGAN (AKTUAL)
    # -------------------------------------------------------------------------
    ws_act = wb.create_sheet(title="Realisasi Lapangan (Aktual)")
    ws_act.views.sheetView[0].showGridLines = True

    ws_act.append([f"LOG REALISASI & BIAYA AKTUAL LAPANGAN - {project.code}"])
    ws_act["A1"].font = title_font
    ws_act.append([])

    # 1. Realisasi Material
    ws_act.append(["1. REALISASI PEMAKAIAN MATERIAL / BAHAN BAKU"])
    ws_act.cell(row=ws_act.max_row, column=1).font = section_title_font
    ws_act.cell(row=ws_act.max_row, column=1).fill = section_title_fill
    ws_act.append(["No", "Tanggal", "Item Material", "Ref. BOM", "Status", "Kuantitas", "Satuan", "Harga Satuan", "Total Biaya", "Catatan / Keterangan"])
    hdr_row = ws_act.max_row
    for c in range(1, 11):
        cell = ws_act.cell(row=hdr_row, column=c)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c in [1, 2, 5, 7] else ("right" if c in [6, 8, 9] else "left"))

    no = 1
    if project.bom_realizations.exists():
        for r in project.bom_realizations.all().order_by("date", "created_at"):
            ref_bom = r.bom_item.name if r.bom_item else "-"
            status_lbl = "Substitusi" if r.is_substitute else "Standar"
            ws_act.append([
                no,
                str(r.date),
                r.item_name,
                ref_bom,
                status_lbl,
                float(r.qty),
                r.unit,
                float(r.unit_cost),
                float(r.total_cost),
                r.notes or "-"
            ])
            curr_row = ws_act.max_row
            for c in range(1, 11):
                cell = ws_act.cell(row=curr_row, column=c)
                cell.border = thin_border
                if c == 6: cell.number_format = qty_fmt; cell.alignment = Alignment(horizontal="right")
                elif c in [8, 9]: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")
                elif c in [1, 2, 5, 7]: cell.alignment = Alignment(horizontal="center")
            no += 1
    else:
        for m in project.bom_items.filter(item_type="material", act_qty__gt=0):
            ws_act.append([
                no,
                str(project.start_date or "-"),
                m.name,
                m.name,
                "Standar",
                float(m.act_qty),
                m.unit,
                float(m.act_unit_cost),
                float(m.act_total),
                m.notes or "-"
            ])
            curr_row = ws_act.max_row
            for c in range(1, 11):
                cell = ws_act.cell(row=curr_row, column=c)
                cell.border = thin_border
                if c == 6: cell.number_format = qty_fmt; cell.alignment = Alignment(horizontal="right")
                elif c in [8, 9]: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")
                elif c in [1, 2, 5, 7]: cell.alignment = Alignment(horizontal="center")
            no += 1

    ws_act.append(["", "SUBTOTAL REALISASI MATERIAL", "", "", "", "", "", "", float(project.total_act_material), ""])
    sub_row = ws_act.max_row
    for c in range(1, 11):
        cell = ws_act.cell(row=sub_row, column=c)
        cell.font = total_font
        cell.fill = sub_header_fill
        cell.border = double_bottom_border
        if c == 9: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")

    ws_act.append([])

    # 2. Realisasi Tenaga Kerja
    ws_act.append(["2. REALISASI TENAGA KERJA (LABOR)"])
    ws_act.cell(row=ws_act.max_row, column=1).font = section_title_font
    ws_act.cell(row=ws_act.max_row, column=1).fill = section_title_fill
    ws_act.append(["No", "Tanggal", "Peran / Posisi", "Tipe", "Volume", "Satuan", "Tarif Satuan", "Total Biaya", "Catatan"])
    hdr_row = ws_act.max_row
    for c in range(1, 10):
        cell = ws_act.cell(row=hdr_row, column=c)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c in [1, 2, 4, 6] else ("right" if c in [5, 7, 8] else "left"))

    no = 1
    if project.labor_realizations.exists():
        for lr in project.labor_realizations.all().order_by("date", "created_at"):
            tipe_lbl = "Tambahan" if lr.is_additional else "Rencana"
            ws_act.append([
                no,
                str(lr.date),
                lr.role_name,
                tipe_lbl,
                float(lr.quantity),
                lr.unit,
                float(lr.rate),
                float(lr.total_cost),
                lr.notes or "-"
            ])
            curr_row = ws_act.max_row
            for c in range(1, 10):
                cell = ws_act.cell(row=curr_row, column=c)
                cell.border = thin_border
                if c == 5: cell.number_format = qty_fmt; cell.alignment = Alignment(horizontal="right")
                elif c in [7, 8]: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")
                elif c in [1, 2, 4, 6]: cell.alignment = Alignment(horizontal="center")
            no += 1
    else:
        for l in project.labor_items.filter(act_quantity__gt=0):
            ws_act.append([
                no,
                str(project.start_date or "-"),
                l.role_name,
                "Rencana",
                float(l.act_quantity),
                l.unit,
                float(l.act_rate),
                float(l.act_total),
                "-"
            ])
            curr_row = ws_act.max_row
            for c in range(1, 10):
                cell = ws_act.cell(row=curr_row, column=c)
                cell.border = thin_border
                if c == 5: cell.number_format = qty_fmt; cell.alignment = Alignment(horizontal="right")
                elif c in [7, 8]: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")
                elif c in [1, 2, 4, 6]: cell.alignment = Alignment(horizontal="center")
            no += 1

    ws_act.append(["", "SUBTOTAL REALISASI TENAGA KERJA", "", "", "", "", "", float(project.total_act_labor), ""])
    sub_row = ws_act.max_row
    for c in range(1, 10):
        cell = ws_act.cell(row=sub_row, column=c)
        cell.font = total_font
        cell.fill = sub_header_fill
        cell.border = double_bottom_border
        if c == 8: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")

    ws_act.append([])

    # 3. Realisasi Barang Jadi
    ws_act.append(["3. REALISASI PEMAKAIAN BARANG JADI (FINISHED GOODS)"])
    ws_act.cell(row=ws_act.max_row, column=1).font = section_title_font
    ws_act.cell(row=ws_act.max_row, column=1).fill = section_title_fill
    ws_act.append(["No", "Tanggal", "SKU", "Nama Barang Jadi", "Kuantitas", "Satuan", "Harga Satuan", "Total Biaya", "Catatan"])
    hdr_row = ws_act.max_row
    for c in range(1, 10):
        cell = ws_act.cell(row=hdr_row, column=c)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c in [1, 2, 3, 6] else ("right" if c in [5, 7, 8] else "left"))

    no = 1
    if project.finished_good_realizations.exists():
        for fgr in project.finished_good_realizations.all().order_by("date", "created_at"):
            ws_act.append([
                no,
                str(fgr.date),
                fgr.finished_good.sku,
                fgr.finished_good.name,
                float(fgr.quantity),
                fgr.finished_good.unit,
                float(fgr.unit_cost),
                float(fgr.total_cost),
                fgr.notes or "-"
            ])
            curr_row = ws_act.max_row
            for c in range(1, 10):
                cell = ws_act.cell(row=curr_row, column=c)
                cell.border = thin_border
                if c == 5: cell.number_format = qty_fmt; cell.alignment = Alignment(horizontal="right")
                elif c in [7, 8]: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")
                elif c in [1, 2, 3, 6]: cell.alignment = Alignment(horizontal="center")
            no += 1
    else:
        for fg in project.finished_good_items.filter(act_qty__gt=0):
            ws_act.append([
                no,
                str(project.start_date or "-"),
                fg.finished_good.sku,
                fg.finished_good.name,
                float(fg.act_qty),
                fg.finished_good.unit,
                float(fg.act_unit_cost),
                float(fg.act_total),
                fg.notes or "-"
            ])
            curr_row = ws_act.max_row
            for c in range(1, 10):
                cell = ws_act.cell(row=curr_row, column=c)
                cell.border = thin_border
                if c == 5: cell.number_format = qty_fmt; cell.alignment = Alignment(horizontal="right")
                elif c in [7, 8]: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")
                elif c in [1, 2, 3, 6]: cell.alignment = Alignment(horizontal="center")
            no += 1

    ws_act.append(["", "SUBTOTAL REALISASI BARANG JADI", "", "", "", "", "", float(project.total_act_finished_goods), ""])
    sub_row = ws_act.max_row
    for c in range(1, 10):
        cell = ws_act.cell(row=sub_row, column=c)
        cell.font = total_font
        cell.fill = sub_header_fill
        cell.border = double_bottom_border
        if c == 8: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")

    ws_act.append([])

    # 4. Realisasi Overhead
    ws_act.append(["4. REALISASI BIAYA OVERHEAD & OPERASIONAL"])
    ws_act.cell(row=ws_act.max_row, column=1).font = section_title_font
    ws_act.cell(row=ws_act.max_row, column=1).fill = section_title_fill
    ws_act.append(["No", "Tanggal", "Deskripsi Biaya Overhead", "Tipe", "Total Biaya", "Catatan"])
    hdr_row = ws_act.max_row
    for c in range(1, 7):
        cell = ws_act.cell(row=hdr_row, column=c)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c in [1, 2, 4] else ("right" if c == 5 else "left"))

    no = 1
    if project.overhead_realizations.exists():
        for orh in project.overhead_realizations.all().order_by("date", "created_at"):
            tipe_lbl = "Tambahan" if orh.is_additional else "Rencana"
            ws_act.append([
                no,
                str(orh.date),
                orh.expense_name,
                tipe_lbl,
                float(orh.cost),
                orh.notes or "-"
            ])
            curr_row = ws_act.max_row
            for c in range(1, 7):
                cell = ws_act.cell(row=curr_row, column=c)
                cell.border = thin_border
                if c == 5: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")
                elif c in [1, 2, 4]: cell.alignment = Alignment(horizontal="center")
            no += 1
    else:
        for o in project.overhead_items.filter(act_cost__gt=0):
            ws_act.append([
                no,
                str(project.start_date or "-"),
                o.name,
                "Rencana",
                float(o.act_cost),
                "-"
            ])
            curr_row = ws_act.max_row
            for c in range(1, 7):
                cell = ws_act.cell(row=curr_row, column=c)
                cell.border = thin_border
                if c == 5: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")
                elif c in [1, 2, 4]: cell.alignment = Alignment(horizontal="center")
            no += 1

    ws_act.append(["", "SUBTOTAL REALISASI OVERHEAD", "", "", float(project.total_act_overhead), ""])
    sub_row = ws_act.max_row
    for c in range(1, 7):
        cell = ws_act.cell(row=sub_row, column=c)
        cell.font = total_font
        cell.fill = sub_header_fill
        cell.border = double_bottom_border
        if c == 5: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")

    ws_act.append([])
    ws_act.append(["GRAND TOTAL REALISASI HPP PROJECT", "", "", "", "", "", "", "", float(project.total_hpp_actual), ""])
    grand_row = ws_act.max_row
    for c in range(1, 11):
        cell = ws_act.cell(row=grand_row, column=c)
        cell.font = Font(name=font_family, size=11, bold=True, color="1E1B4B")
        cell.fill = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
        cell.border = double_bottom_border
        if c == 9: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")

    autofit(ws_act)

    # -------------------------------------------------------------------------
    # SHEET 4: ANALISIS VARIANS (EST VS REAL)
    # -------------------------------------------------------------------------
    ws_var = wb.create_sheet(title="Analisis Varians")
    ws_var.views.sheetView[0].showGridLines = True

    ws_var.append([f"ANALISIS DEVIASI & VARIANS (ESTIMASI VS REALISASI) - {project.code}"])
    ws_var["A1"].font = title_font
    ws_var.append([])

    # Table Header
    ws_var.append([
        "Komponen / Item", "Satuan",
        "Est. Qty", "Est. Total",
        "Real. Qty", "Real. Total",
        "Deviasi Biaya", "Deviasi (%)", "Status Evaluasi"
    ])
    hdr_row = ws_var.max_row
    for c in range(1, 10):
        cell = ws_var.cell(row=hdr_row, column=c)
        cell.font = tbl_header_font
        cell.fill = tbl_header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center" if c in [2, 9] else ("right" if c in [3, 4, 5, 6, 7, 8] else "left"))

    def append_variance_row(ws, name, unit, est_q, est_tot, act_q, act_tot):
        diff = act_tot - est_tot
        ratio = (diff / est_tot) if est_tot > 0 else (0.0 if diff == 0 else 1.0)
        
        if diff < 0:
            status_text, b_font, b_fill = "HEMAT (SAVING)", badge_saving_font, badge_saving_fill
        elif diff > 0:
            status_text, b_font, b_fill = "LEBIH (OVER)", badge_over_font, badge_over_fill
        else:
            status_text, b_font, b_fill = "SESUAI", badge_neutral_font, badge_neutral_fill

        ws.append([name, unit, est_q, est_tot, act_q, act_tot, diff, ratio, status_text])
        r_idx = ws.max_row
        for c in range(1, 10):
            cell = ws.cell(row=r_idx, column=c)
            cell.border = thin_border
            if c in [3, 5]: cell.number_format = qty_fmt; cell.alignment = Alignment(horizontal="right")
            elif c in [4, 6, 7]: cell.number_format = currency_fmt; cell.alignment = Alignment(horizontal="right")
            elif c == 8: cell.number_format = percent_fmt; cell.alignment = Alignment(horizontal="right")
            elif c == 2: cell.alignment = Alignment(horizontal="center")
            elif c == 9: cell.font = b_font; cell.fill = b_fill; cell.alignment = Alignment(horizontal="center")

    # Section Material
    ws_var.append(["A. MATERIAL & BAHAN BAKU", "", "", "", "", "", "", "", ""])
    format_row(ws_var, ws_var.max_row, font=sub_header_font, fill=sub_header_fill, border=thin_border)
    for m in project.bom_items.filter(item_type="material"):
        act_q = float(m.total_realized_qty) if m.realizations.exists() else float(m.act_qty)
        act_t = float(m.total_realized_cost) if m.realizations.exists() else float(m.act_total)
        append_variance_row(ws_var, m.name, m.unit, float(m.est_qty), float(m.est_total), act_q, act_t)

    # Substitusi Material jika ada
    substitutes = project.bom_realizations.filter(Q(is_substitute=True) | Q(bom_item__isnull=True))
    for s in substitutes:
        append_variance_row(ws_var, f"[Substitusi] {s.item_name}", s.unit, 0.0, 0.0, float(s.qty), float(s.total_cost))

    # Section Labor
    ws_var.append(["B. TENAGA KERJA (LABOR)", "", "", "", "", "", "", "", ""])
    format_row(ws_var, ws_var.max_row, font=sub_header_font, fill=sub_header_fill, border=thin_border)
    for l in project.labor_items.all():
        act_q = sum(float(r.quantity) for r in l.realizations.all()) if l.realizations.exists() else float(l.act_quantity)
        act_t = sum(float(r.total_cost) for r in l.realizations.all()) if l.realizations.exists() else float(l.act_total)
        append_variance_row(ws_var, l.role_name, l.unit, float(l.est_quantity), float(l.est_total), act_q, act_t)

    # Section Finished Goods
    ws_var.append(["C. BARANG JADI (FINISHED GOODS)", "", "", "", "", "", "", "", ""])
    format_row(ws_var, ws_var.max_row, font=sub_header_font, fill=sub_header_fill, border=thin_border)
    for fg in project.finished_good_items.all():
        act_q = sum(float(r.quantity) for r in fg.realizations.all()) if fg.realizations.exists() else float(fg.act_qty)
        act_t = sum(float(r.total_cost) for r in fg.realizations.all()) if fg.realizations.exists() else float(fg.act_total)
        append_variance_row(ws_var, f"[{fg.finished_good.sku}] {fg.finished_good.name}", fg.finished_good.unit, float(fg.est_qty), float(fg.est_total), act_q, act_t)

    # Section Overhead
    ws_var.append(["D. OVERHEAD & BIAYA LAINNYA", "", "", "", "", "", "", "", ""])
    format_row(ws_var, ws_var.max_row, font=sub_header_font, fill=sub_header_fill, border=thin_border)
    for o in project.overhead_items.all():
        act_t = sum(float(r.cost) for r in o.realizations.all()) if o.realizations.exists() else float(o.act_cost)
        append_variance_row(ws_var, o.name, "ls", 1.0, float(o.est_cost), 1.0, act_t)

    # Summary Row
    ws_var.append([])
    append_variance_row(
        ws_var,
        "TOTAL KESELURUHAN HPP PROJECT",
        "-",
        0.0, float(project.total_hpp_estimated),
        0.0, float(project.total_hpp_actual)
    )
    r_last = ws_var.max_row
    for c in range(1, 10):
        cell = ws_var.cell(row=r_last, column=c)
        cell.font = Font(name=font_family, size=11, bold=True, color="1E1B4B")
        cell.fill = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
        cell.border = double_bottom_border

    autofit(ws_var)

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="Laporan_HPP_{project.code}.xlsx"'
    wb.save(response)
    return response


def print_project_pdf(request, uuid):
    project = get_object_or_404(Project, uuid=uuid)
    return render(request, "hpp/project_print.html", {
        "project": project,
        "materials": project.bom_items.all().order_by("parent__id", "id"),
        "labors": project.labor_items.all().order_by("id"),
        "overheads": project.overhead_items.all().order_by("id"),
        "finished_goods": project.finished_good_items.all().order_by("id"),
        "bom_realizations": project.bom_realizations.all().order_by("date", "created_at"),
        "labor_realizations": project.labor_realizations.all().order_by("date", "created_at"),
        "overhead_realizations": project.overhead_realizations.all().order_by("date", "created_at"),
        "finished_good_realizations": project.finished_good_realizations.all().order_by("date", "created_at"),
        "now": datetime.now(),
    })


def finished_good_update(request, uuid):
    fg = get_object_or_404(FinishedGood, uuid=uuid)
    if request.method == "POST":
        fg.sku = request.POST.get("sku", fg.sku).strip()
        fg.name = request.POST.get("name", fg.name).strip()
        fg.unit = request.POST.get("unit", fg.unit).strip()
        fg.standard_cost = Decimal(request.POST.get("standard_cost") or 0)
        fg.notes = request.POST.get("notes", fg.notes)
        fg.save()
    return redirect("stock_report")


def finished_good_delete(request, uuid):
    fg = get_object_or_404(FinishedGood, uuid=uuid)
    if request.method == "POST":
        has_mutations = fg.mutations.exists()
        has_projects = fg.project_usages.exists()

        if has_mutations or has_projects:
            # Tidak boleh dihapus jika ada transaksi / relasi project
            from django.contrib import messages
            messages.error(request, f"Barang Jadi '{fg.name}' tidak dapat dihapus karena sudah memiliki riwayat mutasi / transaksi project.")
        else:
            fg.delete()
            from django.contrib import messages
            messages.success(request, f"Barang Jadi '{fg.name}' berhasil dihapus.")
    return redirect("stock_report")


# =========================================================================
# ABSENSI & KARYAWAN VIEWS
# =========================================================================
from datetime import date, datetime


def attendance_list(request):
    """
    Halaman Rekap Absensi Karyawan
    """
    selected_date_str = request.GET.get("date", "").strip()
    status_filter = request.GET.get("status", "").strip()
    absence_type_filter = request.GET.get("absence_type", "").strip()
    query = request.GET.get("q", "").strip()

    if selected_date_str:
        try:
            filter_date = datetime.strptime(selected_date_str, "%Y-%m-%d").date()
        except ValueError:
            filter_date = date.today()
    else:
        filter_date = date.today()

    attendances = Attendance.objects.filter(date=filter_date).select_related("employee", "absence_type")
    if status_filter:
        if status_filter == "HADIR":
            attendances = attendances.filter(status="HADIR")
        elif status_filter == "TIDAK_HADIR":
            attendances = attendances.exclude(status="HADIR")
        else:
            attendances = attendances.filter(status=status_filter)
    if absence_type_filter:
        attendances = attendances.filter(absence_type__uuid=absence_type_filter)
    if query:
        attendances = attendances.filter(Q(employee__name__icontains=query) | Q(employee__nik__icontains=query))

    employees = Employee.objects.filter(is_active=True).order_by("name")
    absence_types = AbsenceType.objects.filter(is_active=True).order_by("category", "code")

    # Summary KPI
    all_day_records = Attendance.objects.filter(date=filter_date).select_related("absence_type")
    total_hadir = all_day_records.filter(status="HADIR").count()
    total_ijin = all_day_records.filter(Q(status="IJIN") | Q(absence_type__category="PERMIT")).count()
    total_sakit = all_day_records.filter(Q(status="SAKIT") | Q(absence_type__category="SICK")).count()
    total_cuti = all_day_records.filter(absence_type__category="LEAVE").count()
    total_dinas = all_day_records.filter(absence_type__category="OFFICIAL_TRAVEL").count()
    total_alpha = all_day_records.filter(Q(status="ALPHA") | Q(absence_type__category="ABSENT")).count()
    total_tidak_hadir = all_day_records.exclude(status="HADIR").count()
    total_ot_hours = sum(a.overtime_hours for a in all_day_records)
    total_wage_day = sum(a.wage for a in all_day_records)

    return render(request, "hpp/attendance_list.html", {
        "attendances": attendances,
        "employees": employees,
        "absence_types": absence_types,
        "selected_date": filter_date.strftime("%Y-%m-%d"),
        "status_filter": status_filter,
        "absence_type_filter": absence_type_filter,
        "query": query,
        "status_choices": Attendance.STATUS_CHOICES,
        "total_hadir": total_hadir,
        "total_ijin": total_ijin,
        "total_sakit": total_sakit,
        "total_cuti": total_cuti,
        "total_dinas": total_dinas,
        "total_alpha": total_alpha,
        "total_tidak_hadir": total_tidak_hadir,
        "total_ot_hours": total_ot_hours,
        "total_wage_day": total_wage_day,
    })


def attendance_create(request):
    if request.method == "POST":
        emp_uuid = request.POST.get("employee_uuid")
        att_date_str = request.POST.get("date") or str(date.today())
        status = request.POST.get("status", "HADIR")
        absence_type_uuid = request.POST.get("absence_type_uuid", "").strip()
        check_in = request.POST.get("check_in") or None
        check_out = request.POST.get("check_out") or None
        overtime_hours = Decimal(request.POST.get("overtime_hours") or 0)
        notes = request.POST.get("notes", "").strip()

        employee = get_object_or_404(Employee, uuid=emp_uuid)
        try:
            att_date = datetime.strptime(att_date_str, "%Y-%m-%d").date()
        except ValueError:
            att_date = date.today()

        absence_type = None
        if status == "HADIR":
            absence_type = None
        else:
            if absence_type_uuid:
                absence_type = AbsenceType.objects.filter(uuid=absence_type_uuid).first()
                status = "TIDAK_HADIR"
            elif status in ["IJIN", "SAKIT", "ALPHA"]:
                pass
            else:
                status = "TIDAK_HADIR"

            # Reset jam jika bukan dinas luar
            if not absence_type or absence_type.category != "OFFICIAL_TRAVEL":
                check_in = None
                check_out = None
                overtime_hours = Decimal(0)

        att, _ = Attendance.objects.update_or_create(
            employee=employee,
            date=att_date,
            defaults={
                "status": status,
                "absence_type": absence_type,
                "check_in": check_in,
                "check_out": check_out,
                "overtime_hours": overtime_hours,
                "notes": notes,
            }
        )
        messages.success(request, f"Absensi {employee.name} tanggal {att_date} ({att.display_status_label}) berhasil disimpan.")
        return redirect(f"/attendance/?date={att_date_str}")
    return redirect("attendance_list")


def attendance_update(request, uuid):
    att = get_object_or_404(Attendance, uuid=uuid)
    if request.method == "POST":
        status = request.POST.get("status", att.status)
        absence_type_uuid = request.POST.get("absence_type_uuid", "").strip()
        check_in = request.POST.get("check_in") or None
        check_out = request.POST.get("check_out") or None
        overtime_hours = Decimal(request.POST.get("overtime_hours") or 0)
        notes = request.POST.get("notes", "").strip()

        if status == "HADIR":
            att.status = "HADIR"
            att.absence_type = None
            att.check_in = check_in
            att.check_out = check_out
            att.overtime_hours = overtime_hours
        else:
            if absence_type_uuid:
                att.absence_type = AbsenceType.objects.filter(uuid=absence_type_uuid).first()
                att.status = "TIDAK_HADIR"
            elif status in ["IJIN", "SAKIT", "ALPHA"]:
                att.status = status
            else:
                att.status = "TIDAK_HADIR"

            if not att.absence_type or att.absence_type.category != "OFFICIAL_TRAVEL":
                att.check_in = None
                att.check_out = None
                att.overtime_hours = Decimal(0)

        att.notes = notes
        att.save()
        messages.success(request, f"Absensi {att.employee.name} ({att.display_status_label}) berhasil diperbarui.")
    return redirect(f"/attendance/?date={att.date.strftime('%Y-%m-%d')}")


def attendance_delete(request, uuid):
    att = get_object_or_404(Attendance, uuid=uuid)
    att_date_str = att.date.strftime("%Y-%m-%d")
    if request.method == "POST":
        att.delete()
        messages.success(request, f"Data absensi berhasil dihapus.")
    return redirect(f"/attendance/?date={att_date_str}")


def employee_list(request):
    """
    Daftar Master Karyawan
    """
    employees = Employee.objects.all().order_by("name")
    return render(request, "hpp/employee_list.html", {
        "employees": employees,
    })


def employee_create(request):
    if request.method == "POST":
        nik = request.POST.get("nik", "").strip()
        name = request.POST.get("name", "").strip()
        position = request.POST.get("position", "Tukang").strip()
        daily_rate = Decimal(request.POST.get("daily_rate") or 0)
        overtime_rate_per_hour = Decimal(request.POST.get("overtime_rate_per_hour") or 0)
        notes = request.POST.get("notes", "")

        Employee.objects.create(
            nik=nik,
            name=name,
            position=position,
            daily_rate=daily_rate,
            overtime_rate_per_hour=overtime_rate_per_hour,
            notes=notes,
        )
        messages.success(request, f"Karyawan {name} berhasil ditambahkan.")
    return redirect("employee_list")


def employee_update(request, uuid):
    emp = get_object_or_404(Employee, uuid=uuid)
    if request.method == "POST":
        emp.nik = request.POST.get("nik", emp.nik).strip()
        emp.name = request.POST.get("name", emp.name).strip()
        emp.position = request.POST.get("position", emp.position).strip()
        emp.daily_rate = Decimal(request.POST.get("daily_rate") or 0)
        emp.overtime_rate_per_hour = Decimal(request.POST.get("overtime_rate_per_hour") or 0)
        emp.is_active = request.POST.get("is_active") == "on"
        emp.notes = request.POST.get("notes", emp.notes)
        emp.save()
        messages.success(request, f"Data karyawan {emp.name} berhasil diperbarui.")
    return redirect("employee_list")


def employee_delete(request, uuid):
    emp = get_object_or_404(Employee, uuid=uuid)
    if request.method == "POST":
        if emp.attendances.exists():
            messages.error(request, f"Karyawan {emp.name} tidak dapat dihapus karena memiliki riwayat absensi. Silakan nonaktifkan statusnya.")
        else:
            emp.delete()
            messages.success(request, f"Karyawan {emp.name} berhasil dihapus.")
    return redirect("employee_list")


# =========================================================================
# KINERJA KARYAWAN (WORK LOGS) VIEWS
# =========================================================================
STANDARD_ACTIVITY_CATEGORIES = [
    "Fabrikasi & Pemotongan Kayu / Besi",
    "Perakitan & Konstruksi (Assembly)",
    "Finishing, Dempul & Pengamplasan",
    "Pengecatan / Coating / Melamik",
    "Pemasangan Hardware & Aksesoris",
    "Packing & Quality Control (QC)",
    "Pemasangan di Lokasi Proyek (Site Installation)",
    "Maintenance & Pemeliharaan Alat",
    "Pekerjaan Umum / Workshop",
]


def work_log_list(request):
    """
    Daftar Log Kinerja & Kegiatan Harian Karyawan
    Dilengkapi filter proyek, pencarian, indikator presensi, dan konversi ke realisasi HPP.
    """
    selected_date_str = request.GET.get("date", "").strip()
    emp_uuid = request.GET.get("employee", "").strip()
    proj_uuid = request.GET.get("project", "").strip()
    status_filter = request.GET.get("status", "").strip()
    query = request.GET.get("q", "").strip()

    logs = EmployeeWorkLog.objects.all().select_related("employee", "project", "labor_realization")

    if selected_date_str:
        try:
            filter_date = datetime.strptime(selected_date_str, "%Y-%m-%d").date()
            logs = logs.filter(date=filter_date)
        except ValueError:
            pass

    if emp_uuid:
        logs = logs.filter(employee__uuid=emp_uuid)

    if proj_uuid:
        logs = logs.filter(project__uuid=proj_uuid)

    if status_filter:
        logs = logs.filter(status=status_filter)

    if query:
        logs = logs.filter(
            Q(task_description__icontains=query) |
            Q(activity_category__icontains=query) |
            Q(employee__name__icontains=query) |
            Q(employee__nik__icontains=query) |
            Q(project__name__icontains=query) |
            Q(project__code__icontains=query) |
            Q(obstacles__icontains=query)
        )

    employees = Employee.objects.filter(is_active=True).order_by("name")
    projects = Project.objects.prefetch_related("labor_items").order_by("-created_at")
    projects_data = []
    for p in projects:
        labor_items_data = []
        for li in p.labor_items.all():
            labor_items_data.append({
                "uuid": str(li.uuid),
                "role_name": li.role_name,
                "unit": li.unit,
                "est_quantity": float(li.est_quantity),
                "est_rate": float(li.est_rate),
                "label": f"{li.role_name} (Budget: {li.est_quantity:g} {li.unit} @ Rp {li.est_rate:,.0f})",
            })
        projects_data.append({
            "uuid": str(p.uuid),
            "code": p.code,
            "name": p.name,
            "status": p.status,
            "status_display": p.get_status_display(),
            "customer": p.customer_name or (p.customer.name if p.customer else ""),
            "label": f"[{p.code}] {p.name}",
            "is_completed": p.status == "completed",
            "labor_items": labor_items_data,
        })
    projects_json = json.dumps(projects_data)

    total_logs = logs.count()
    total_hours = sum(l.hours_spent for l in logs)
    completed_count = logs.filter(status="completed").count()

    # Pre-fetch attendance for all (employee_id, date) pairs in logs to avoid N+1 queries
    log_emp_ids = set(l.employee_id for l in logs)
    log_dates = set(l.date for l in logs)
    attendance_map = {}
    if log_emp_ids and log_dates:
        attendances = Attendance.objects.filter(
            employee_id__in=log_emp_ids,
            date__in=log_dates
        ).select_related("absence_type")
        for att in attendances:
            attendance_map[(att.employee_id, att.date)] = att

    for l in logs:
        l.attendance = attendance_map.get((l.employee_id, l.date))

    return render(request, "hpp/work_log_list.html", {
        "logs": logs,
        "employees": employees,
        "projects": projects,
        "projects_json": projects_json,
        "selected_date": selected_date_str,
        "selected_emp": emp_uuid,
        "selected_proj": proj_uuid,
        "selected_status": status_filter,
        "query": query,
        "status_choices": EmployeeWorkLog.STATUS_CHOICES,
        "rating_choices": EmployeeWorkLog.RATING_CHOICES,
        "standard_categories": STANDARD_ACTIVITY_CATEGORIES,
        "total_logs": total_logs,
        "total_hours": total_hours,
        "completed_count": completed_count,
    })


def work_log_create(request):
    if request.method == "POST":
        emp_uuid = request.POST.get("employee_uuid")
        proj_uuid = request.POST.get("project_uuid") or None
        date_str = request.POST.get("date") or str(date.today())
        activity_category = request.POST.get("activity_category", "Fabrikasi & Produksi").strip()
        task_description = request.POST.get("task_description", "").strip()
        output_qty = Decimal(request.POST.get("output_qty") or 0)
        output_unit = request.POST.get("output_unit", "unit").strip()
        hours_spent = Decimal(request.POST.get("hours_spent") or 8)
        status = request.POST.get("status", "completed")
        obstacles = request.POST.get("obstacles", "").strip()
        supervisor_rating = int(request.POST.get("supervisor_rating") or 4)

        employee = get_object_or_404(Employee, uuid=emp_uuid)
        project = Project.objects.filter(uuid=proj_uuid).first() if proj_uuid else None
        log_date = datetime.strptime(date_str, "%Y-%m-%d").date()

        # Cross-validation with Attendance
        att = Attendance.objects.filter(employee=employee, date=log_date).select_related("absence_type").first()
        if att and att.status != "HADIR":
            status_label = att.display_status_label
            messages.warning(
                request,
                f"Pemberitahuan Presensi: Karyawan {employee.name} tercatat '{status_label}' "
                f"pada tanggal {log_date}. Log kegiatan tetap disimpan sebagai catatan operasional."
            )

        EmployeeWorkLog.objects.create(
            employee=employee,
            project=project,
            date=log_date,
            activity_category=activity_category,
            task_description=task_description,
            output_qty=output_qty,
            output_unit=output_unit,
            hours_spent=hours_spent,
            status=status,
            obstacles=obstacles,
            supervisor_rating=supervisor_rating,
        )
        messages.success(request, f"Log kegiatan {employee.name} berhasil disimpan.")
    return redirect("work_log_list")


def work_log_update(request, uuid):
    log = get_object_or_404(EmployeeWorkLog, uuid=uuid)
    if request.method == "POST":
        proj_uuid = request.POST.get("project_uuid") or None
        date_str = request.POST.get("date") or str(log.date)
        log.date = datetime.strptime(date_str, "%Y-%m-%d").date()
        log.project = Project.objects.filter(uuid=proj_uuid).first() if proj_uuid else None
        log.activity_category = request.POST.get("activity_category", log.activity_category).strip()
        log.task_description = request.POST.get("task_description", log.task_description).strip()
        log.output_qty = Decimal(request.POST.get("output_qty") or 0)
        log.output_unit = request.POST.get("output_unit", log.output_unit).strip()
        log.hours_spent = Decimal(request.POST.get("hours_spent") or 8)
        log.status = request.POST.get("status", log.status)
        log.obstacles = request.POST.get("obstacles", log.obstacles).strip()
        log.supervisor_rating = int(request.POST.get("supervisor_rating") or log.supervisor_rating)

        # Cross-validation with Attendance
        att = Attendance.objects.filter(employee=log.employee, date=log.date).select_related("absence_type").first()
        if att and att.status != "HADIR":
            messages.warning(
                request,
                f"Pemberitahuan Presensi: Karyawan {log.employee.name} tercatat '{att.display_status_label}' pada tanggal {log.date}."
            )

        log.save()
        messages.success(request, f"Log kegiatan {log.employee.name} berhasil diperbarui.")
    return redirect("work_log_list")


def work_log_delete(request, uuid):
    log = get_object_or_404(EmployeeWorkLog, uuid=uuid)
    if request.method == "POST":
        # If linked to labor realization, delete realization as well or warn
        if log.labor_realization:
            if log.project and log.project.status == "completed":
                messages.error(request, f"Proyek '{log.project.name}' sudah selesai (Terkunci). Log tidak dapat dihapus.")
                return redirect("work_log_list")
            log.labor_realization.delete()

        log.delete()
        messages.success(request, f"Log kegiatan berhasil dihapus.")
    return redirect("work_log_list")


@transaction.atomic
def work_log_post_to_realization(request, uuid):
    """
    Posting Jam Kerja dari Log Kinerja ke Realisasi Biaya Tenaga Kerja Proyek HPP
    """
    log = get_object_or_404(EmployeeWorkLog, uuid=uuid)
    if request.method == "POST":
        if not log.project:
            messages.error(request, "Log kinerja tidak dapat dikonversi ke HPP karena tidak terhubung ke proyek manapun.")
            return redirect("work_log_list")

        if log.project.status == "completed":
            messages.error(request, f"Proyek '{log.project.name}' sudah berstatus Selesai (Audit Locked). Penambahan realisasi biaya ditolak.")
            return redirect("work_log_list")

        if log.labor_realization:
            messages.warning(request, f"Log kinerja ini sudah pernah dikonversi ke Realisasi Tenaga Kerja Proyek ({log.project.code}).")
            return redirect("work_log_list")

        custom_rate_str = request.POST.get("rate", "").strip()
        if custom_rate_str:
            try:
                rate = Decimal(custom_rate_str)
            except Exception:
                rate = log.suggested_hourly_rate
        else:
            rate = log.suggested_hourly_rate

        labor_item_uuid = request.POST.get("labor_item_uuid", "").strip()
        project_labor = None
        if labor_item_uuid:
            project_labor = ProjectLabor.objects.filter(uuid=labor_item_uuid, project=log.project).first()

        role_name = request.POST.get("role_name", "").strip()
        if not role_name:
            if project_labor:
                role_name = f"{project_labor.role_name} ({log.employee.name})"
            else:
                role_name = f"{log.employee.position} ({log.employee.name})"

        is_additional_param = request.POST.get("is_additional")
        if is_additional_param is not None:
            is_additional = is_additional_param in ["on", "true", "1"]
        else:
            is_additional = project_labor is None

        notes = request.POST.get("notes", "").strip() or f"Auto-post dari Log Kinerja: {log.task_description[:80]}"

        realization = LaborRealization.objects.create(
            project=log.project,
            project_labor=project_labor,
            date=log.date,
            role_name=role_name,
            unit="jam",
            quantity=log.hours_spent,
            rate=rate,
            is_additional=is_additional,
            notes=notes,
        )

        log.labor_realization = realization
        log.save(update_fields=["labor_realization"])

        messages.success(
            request,
            f"Log kinerja {log.employee.name} berhasil diposting ke Realisasi Tenaga Kerja Proyek "
            f"'{log.project.name}' (Biaya: Rp {realization.total_cost:,.0f})."
        )
    return redirect("work_log_list")


@transaction.atomic
def work_log_unpost_from_realization(request, uuid):
    """
    Batalkan Konversi Log Kinerja dari Realisasi Tenaga Kerja Proyek
    """
    log = get_object_or_404(EmployeeWorkLog, uuid=uuid)
    if request.method == "POST":
        if not log.labor_realization:
            messages.info(request, "Log kinerja ini tidak sedang terhubung ke realisasi biaya proyek.")
            return redirect("work_log_list")

        if log.project and log.project.status == "completed":
            messages.error(request, f"Proyek '{log.project.name}' sudah selesai (Audit Locked). Realisasi tidak dapat dibatalkan.")
            return redirect("work_log_list")

        realization = log.labor_realization
        log.labor_realization = None
        log.save(update_fields=["labor_realization"])
        realization.delete()

        messages.success(request, f"Realisasi tenaga kerja proyek untuk log {log.employee.name} berhasil dibatalkan.")
    return redirect("work_log_list")


def work_log_export_excel(request):
    """
    Export Log Kinerja & Kegiatan Karyawan ke Format Excel (.xlsx)
    """
    selected_date_str = request.GET.get("date", "").strip()
    emp_uuid = request.GET.get("employee", "").strip()
    proj_uuid = request.GET.get("project", "").strip()
    status_filter = request.GET.get("status", "").strip()
    query = request.GET.get("q", "").strip()

    logs = EmployeeWorkLog.objects.all().select_related("employee", "project", "labor_realization")

    if selected_date_str:
        try:
            filter_date = datetime.strptime(selected_date_str, "%Y-%m-%d").date()
            logs = logs.filter(date=filter_date)
        except ValueError:
            pass

    if emp_uuid:
        logs = logs.filter(employee__uuid=emp_uuid)

    if proj_uuid:
        logs = logs.filter(project__uuid=proj_uuid)

    if status_filter:
        logs = logs.filter(status=status_filter)

    if query:
        logs = logs.filter(
            Q(task_description__icontains=query) |
            Q(activity_category__icontains=query) |
            Q(employee__name__icontains=query) |
            Q(employee__nik__icontains=query) |
            Q(project__name__icontains=query) |
            Q(project__code__icontains=query) |
            Q(obstacles__icontains=query)
        )

    logs = logs.order_by("-date", "employee__name")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Log Kinerja SDM"

    title_font = Font(name="Calibri", size=14, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    regular_font = Font(name="Calibri", size=10)
    bold_font = Font(name="Calibri", size=10, bold=True)
    thin_side = Side(border_style="thin", color="CBD5E1")
    thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    ws["A1"] = "LAPORAN LOG KINERJA & KEGIATAN KARYAWAN"
    ws["A1"].font = title_font
    ws["A2"] = f"Dicetak pada: {timezone.now().strftime('%d/%m/%Y %H:%M:%S')} | Total Catatan: {logs.count()}"
    ws["A2"].font = subtitle_font

    headers = [
        "No", "Tanggal", "NIK", "Nama Karyawan", "Posisi",
        "Proyek Terkait", "Kategori Aktivitas", "Deskripsi Pekerjaan",
        "Output Qty", "Satuan", "Durasi (Jam)", "Status",
        "Rating Mandor", "Kendala Lapangan", "Status HPP Proyek"
    ]

    for col_num, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    row_idx = 5
    total_hours = Decimal(0)
    for idx, l in enumerate(logs, 1):
        total_hours += l.hours_spent
        proj_str = f"[{l.project.code}] {l.project.name}" if l.project else "Non-Project / Workshop"
        hpp_status = f"Terkonversi (Rp {l.labor_realization.total_cost:,.0f})" if l.labor_realization else ("Tersedia" if l.project else "-")

        ws.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=l.date.strftime("%d/%m/%Y")).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=3, value=l.employee.nik).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=4, value=l.employee.name)
        ws.cell(row=row_idx, column=5, value=l.employee.position)
        ws.cell(row=row_idx, column=6, value=proj_str)
        ws.cell(row=row_idx, column=7, value=l.activity_category)
        ws.cell(row=row_idx, column=8, value=l.task_description)
        ws.cell(row=row_idx, column=9, value=float(l.output_qty)).alignment = Alignment(horizontal="right")
        ws.cell(row=row_idx, column=10, value=l.output_unit).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=11, value=float(l.hours_spent)).alignment = Alignment(horizontal="right")
        ws.cell(row=row_idx, column=12, value=l.get_status_display()).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=13, value=f"{l.supervisor_rating}/5").alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=14, value=l.obstacles or "-")
        ws.cell(row=row_idx, column=15, value=hpp_status).alignment = Alignment(horizontal="center")

        for col_num in range(1, 16):
            c = ws.cell(row=row_idx, column=col_num)
            c.font = regular_font
            c.border = thin_border

        row_idx += 1

    ws.cell(row=row_idx, column=1, value="TOTAL JAM KERJA")
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=10)
    ws.cell(row=row_idx, column=1).font = bold_font
    ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="right")

    tot_cell = ws.cell(row=row_idx, column=11, value=float(total_hours))
    tot_cell.font = bold_font
    tot_cell.alignment = Alignment(horizontal="right")

    for col_num in range(1, 16):
        ws.cell(row=row_idx, column=col_num).border = thin_border

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 11)

    ws.column_dimensions["H"].width = 35
    ws.column_dimensions["F"].width = 28
    ws.column_dimensions["D"].width = 22

    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    timestamp = timezone.now().strftime("%Y%m%d_%H%M%S")
    response["Content-Disposition"] = f'attachment; filename="Log_Kinerja_SDM_{timestamp}.xlsx"'
    wb.save(response)
    return response



# =========================================================================
# MUTASI BARANG JADI & KARTU STOCK DEDICATED VIEWS
# =========================================================================
from . import views_finished_good

def stock_mutation_list(request):
    """Delegasi ke modul views_finished_good"""
    return views_finished_good.finished_good_mutation_list(request)


def stock_mutation_general_create(request):
    """Delegasi ke modul views_finished_good"""
    return views_finished_good.finished_good_mutation_create(request)


def stock_card_index(request):
    """
    Menu Kartu Stock Standalone -> Delegasi ke views_finished_good
    """
    return views_finished_good.finished_good_stock_card_index(request)


def finished_good_form_view(request):
    """
    Halaman Form Tambah Barang Jadi
    """
    if request.method == "POST":
        sku = request.POST.get("sku", "").strip()
        name = request.POST.get("name", "").strip()
        unit = request.POST.get("unit", "unit").strip()
        standard_cost = Decimal(request.POST.get("standard_cost") or 0)
        initial_stock = Decimal(request.POST.get("initial_stock") or 0)
        notes = request.POST.get("notes", "")

        if FinishedGood.objects.filter(sku=sku).exists():
            messages.error(request, f"SKU {sku} sudah terdaftar.")
            return render(request, "hpp/finished_good_form.html")

        fg = FinishedGood.objects.create(
            sku=sku,
            name=name,
            unit=unit,
            standard_cost=standard_cost,
            current_stock=initial_stock,
            notes=notes,
        )

        if initial_stock > 0:
            FinishedGoodStockMutation.objects.create(
                finished_good=fg,
                mutation_type="IN",
                quantity=initial_stock,
                balance_after=initial_stock,
                reference_no=f"INIT-{fg.sku}",
                notes="Saldo stock awal"
            )
        messages.success(request, f"Barang Jadi '{name}' berhasil ditambahkan.")
        return redirect("stock_report")

    return render(request, "hpp/finished_good_form.html")
